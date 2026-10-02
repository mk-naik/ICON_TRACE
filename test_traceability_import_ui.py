"""
ICON TRACE - the traceability importer on the Production Entry SCREEN, end to
end in a real browser.

    python test_traceability_import_ui.py

test_traceability_import.py pins down the parser and the two apply routes.
This pins down what the SCREEN does with them: the "Upload Excel" tab parses
the file, offers date -> shift -> range(s), and records the ticked ranges -
in backfill mode here, because a fresh test database has planned nothing, so
that is the path that actually creates production.

ui_harness is imported FIRST, on purpose: it fixes the database path the
server and this test both use, before app binds to it.
"""

import io, os, sys, tempfile, traceback, datetime

import ui_harness as H                                       # noqa: E402  (first: sets the DB path)
import store                                                 # noqa: E402
import openpyxl                                              # noqa: E402

TMP = tempfile.mkdtemp(prefix="icontrace_traceui_")
_results = []


def test(name):
    def deco(fn):
        _results.append((name, fn))
        return fn
    return deco


HEADERS = ["Date", "Shift", "Module Wattage", "Serial No.", "Ending No",
           "Quantity", "Special Customer", "Remark", "Cell Make",
           "Cell Invoice/Batch no", "Cell Type"]
BOM = ["LIONSOLAR  25.7% (210*182.2) G12R CELL", "ID20260701", "G12R"]


def make_file(rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TRACEABILITY TEST"
    ws.cell(3, 7, "... for the month TEST - 2026")
    for c, h in enumerate(HEADERS):
        ws.cell(12, 2 + c, h)          # column A empty, like the real file
    r = 13
    for row in rows:
        vals = [row["date"], row["shift"], row["watt"], row["start"], row["end"],
                row["qty"], row["customer"], True] + BOM
        for c, v in enumerate(vals):
            if v is not None:
                ws.cell(r, 2 + c, v)
        r += 1
    path = os.path.join(TMP, "trace.xlsx")
    wb.save(path)
    return path


def S(seq):
    return "ICON625R129014%04d" % seq


DATA = [
    {"date": datetime.datetime(2026, 9, 1), "shift": "A", "watt": "625W",
     "start": S(1), "end": S(7), "qty": 7, "customer": "BOROSIL"},
    {"date": datetime.datetime(2026, 9, 1), "shift": "B", "watt": "625W",
     "start": S(8), "end": S(20), "qty": 13, "customer": "BOROSIL"},
]


@test("Upload Excel -> pick date/shift -> tick a range -> backfill records it: "
      "the cascade drives the real import end to end in the browser")
def t_import_cascade():
    with store.conn() as (cx, cur):          # the file names no incharge: one is chosen
        __import__("db").incharge_add(cur, ["X"])
    path = make_file(DATA)
    with H.browser() as b:
        pg = H.open_page(b, "prodentry")
        # the entry form is collapsed behind "New production entry" until asked for
        pg.evaluate("peToggleForm(true)")
        pg.wait_for_timeout(200)
        # switch to the Upload Excel tab (v4's own seg button)
        pg.click("#v-prodentry .seg button:has-text('Upload Excel')")
        pg.wait_for_timeout(300)
        pg.set_input_files("#peFileIn", path)
        pg.wait_for_selector("#peImpDate", timeout=8000)

        # the cascade knows both dates' one date and both shifts
        dates = pg.eval_on_selector_all("#peImpDate option", "els => els.map(e => e.value)")
        assert dates == ["2026-09-01"], dates
        shifts = pg.eval_on_selector_all("#peImpShift option", "els => els.map(e => e.textContent)")
        assert shifts == ["A", "B"], shifts

        # shift A has one range, 7 modules; it reads "not in system" (nothing planned)
        pg.wait_for_selector("#peImpRanges table tbody tr")
        assert "not in system" in pg.inner_text("#peImpRanges").lower(), pg.inner_text("#peImpRanges")

        # turn on backfill (so the disabled/enabled state flips to ready) and record
        pg.check("#peImpBackfill")
        pg.evaluate("window.__peImpPicker.set('X')")
        pg.wait_for_timeout(200)
        pg.click("#peImpApply")
        pg.wait_for_selector("#peImpResult .note", timeout=8000)
        msg = pg.inner_text("#peImpResult")
        assert "Recorded 1" in msg and "7 modules" in msg, msg
        assert not pg.errors, pg.errors

    # the 7 serials are now really in the database, produced
    with store.conn() as (cx, cur):
        n = store.one(cur, "SELECT COUNT(*) AS n FROM serial WHERE state='produced'")["n"]
        pe = store.one(cur, "SELECT COUNT(*) AS n FROM production_entry")["n"]
        mats = store.one(cur, "SELECT COUNT(*) AS n FROM allocation_material")["n"]
    assert n == 7, n
    assert pe == 1, pe
    assert mats >= 1, "the BOM was captured on the allocation"


@test("Clear form wipes the parsed import too - the cascade, the chosen file "
      "and the result, not just the manual fields")
def t_clear_form_clears_import():
    path = make_file(DATA)
    with H.browser() as b:
        pg = H.open_page(b, "prodentry")
        pg.evaluate("peToggleForm(true)")
        pg.wait_for_timeout(200)
        pg.click("#v-prodentry .seg button:has-text('Upload Excel')")
        pg.set_input_files("#peFileIn", path)
        pg.wait_for_selector("#peImpDate", timeout=8000)
        assert pg.is_visible("#peImpBody"), "cascade is up before clearing"
        pg.evaluate("peClearForm()")
        pg.wait_for_timeout(200)
        assert not pg.is_visible("#peImpBody"), "the parsed cascade is gone after Clear form"
        assert pg.evaluate("window.__peImport") is None, "parsed data dropped"
        assert pg.eval_on_selector("#peFileIn", "e => e.value") == "", "file input reset"
        assert not pg.errors, pg.errors


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    width = max(len(n) for n, _ in _results)
    passed = failed = 0
    try:
        for name, fn in _results:
            try:
                fn()
                print("  PASS  %-*s" % (width, name))
                passed += 1
            except Exception as e:
                print("  FAIL  %-*s  %s" % (width, name, e))
                if "-v" in sys.argv:
                    traceback.print_exc()
                failed += 1
        print("\n%d passed, %d failed" % (passed, failed))
    finally:
        H.cleanup()
        import shutil
        shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(1 if failed else 0)
