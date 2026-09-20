"""
ICON TRACE - tests for exporting the indent list.

    python test_indent_export.py      (needs Playwright + Chromium)

THE RULES THIS FILE DEFENDS

  1. THE INDENT LIST EXPORTS THROUGH THE SAME MECHANISM AS EVERY OTHER
     SCREEN - the rows are read off the screen and sent to /api/export/xlsx.
     Its Export button used to sit outside the table's own wiring and do
     nothing at all.

  2. ONE ROW PER INDENT ITEM, in the columns a spreadsheet can be filtered
     and pivoted on: indent, customer, item, cell type, ordered, dispatched,
     remaining. The screen paints for the eye - a customer shown once per
     indent, an "item 2" note under the number - and a file read by a filter
     needs both filled on every row, or the second item of an indent drops
     out of a customer's total.

  3. WHAT A FILTER HID IS NOT IN THE FILE, as everywhere else.

  4. AN INDENT WITH NO ITEM IS NOT AN ITEM ROW. The screen lists it so it can
     be fixed; the export is of items.
"""

import io, sys, traceback

import ui_harness as H                                       # noqa: E402  (first: sets the DB path)
import db                                                    # noqa: E402
import store                                                 # noqa: E402
import icon_models as models                                 # noqa: E402

APP = H.APP
_results = []


def test(name):
    def deco(fn):
        _results.append((name, fn))
        return fn
    return deco


ITEMS = [i["item_code"] for i in models.all_items()][:3]
A, B = "SEP-10/2026", "SEP-09/2026"
CUST_A, CUST_B = "SAI BABUJI PROJECTS", "SG MEDA"
COLUMNS = ["Indent", "Customer", "Item", "Cell type", "Ordered",
           "Dispatched", "Remaining"]


def item(code):
    return models.get_item(code)


def seed():
    """A: two items for one customer. B: one item for another. C: an indent
    somebody emptied - on the list, not in the export."""
    store.wipe()
    c = APP.app.test_client()
    for no, date, cust, items in (
            (A, "2026-09-10", CUST_A, [(ITEMS[0], 120), (ITEMS[1], 60)]),
            (B, "2026-09-09", CUST_B, [(ITEMS[2], 36)])):
        r = c.post("/api/indent", json={
            "indent_no": no, "indent_date": date, "customer": cust,
            "items": [{"item_code": code, "qty": q} for code, q in items]})
        assert r.get_json().get("ok"), r.get_json()
    with store.conn() as (cx, cur):
        db.insert_indent(cur, {"indent_no": "SEP-08/2026", "indent_date": "2026-09-08",
                               "customer": "SG MEDA", "build_type": "make_to_stock",
                               "status": "open", "form_no": "IS-HO-MRK-FM-03"},
                         [], None, None, "tester")
    return c


def expected(entries):
    return [[no, cust, item(code)["item"], item(code)["cell_type"], str(q), "0", str(q)]
            for no, cust, code, q in entries]


ALL_ROWS = [(A, CUST_A, ITEMS[0], 120), (A, CUST_A, ITEMS[1], 60),
            (B, CUST_B, ITEMS[2], 36)]


def export_request(pg):
    """Click Export and catch what the page sends to the server."""
    with pg.expect_request("**/api/export/xlsx", timeout=8000) as info:
        pg.click("#v-indent .pg-act >> text=Export table")
    return info.value.post_data_json


def open_indent(b):
    pg = H.open_page(b, "indent", wait_ms=900)
    pg.wait_for_selector("#indRows tr")
    pg.wait_for_timeout(300)
    return pg


@test("Export on the indent list reaches the generic xlsx export, and sends "
      "one row per indent item in the seven columns - customer and indent "
      "number filled on the second item too, the empty indent left out")
def t_export_rows():
    seed()
    with H.browser() as b:
        pg = open_indent(b)
        # the screen really does blank the customer on an indent's second
        # item - which is exactly why the file cannot be a raw copy of it
        shown = pg.eval_on_selector_all(
            "#indRows tr", "rs => rs.map(r => r.cells[1].textContent.trim())")
        assert len(shown) == 4, "the emptied indent is not on the list to leave out: %s" % shown
        assert shown.count("") == 1, "screen no longer blanks a repeated customer: %s" % shown

        body = export_request(pg)
        assert len(body["sheets"]) == 1, [s["title"] for s in body["sheets"]]
        sheet = body["sheets"][0]
        assert sheet["columns"] == COLUMNS, sheet["columns"]
        assert sheet["rows"] == expected(ALL_ROWS), sheet["rows"]
        assert not pg.errors, pg.errors


@test("a filtered list exports only what is showing")
def t_export_respects_filter():
    seed()
    with H.browser() as b:
        pg = open_indent(b)
        pg.select_option("#ixCust", label=CUST_B)
        pg.wait_for_timeout(200)
        body = export_request(pg)
        assert body["sheets"][0]["rows"] == expected([ALL_ROWS[2]]), body["sheets"][0]["rows"]


@test("what comes back is a real workbook: the seven headers, the same rows, "
      "quantities as numbers that can be summed")
def t_workbook():
    from openpyxl import load_workbook
    c = seed()
    with H.browser() as b:
        pg = open_indent(b)
        body = export_request(pg)
    r = c.post("/api/export/xlsx", json=body)
    assert r.status_code == 200, r.get_data()[:200]
    ws = load_workbook(io.BytesIO(r.data)).active
    cells = [[cell.value for cell in row] for row in ws.iter_rows()]
    assert cells[0] == COLUMNS, cells[0]
    assert len(cells) == 1 + len(ALL_ROWS), cells
    for got, (no, cust, code, q) in zip(cells[1:], ALL_ROWS):
        assert got[:4] == [no, cust, item(code)["item"], item(code)["cell_type"]], got
        assert got[4:] == [q, 0, q], "quantities are not numbers: %r" % (got[4:],)
    assert sum(row[4] for row in cells[1:]) == 216


@test("the screen itself is unchanged: the second item still shows no "
      "customer, and the item note and status are still painted")
def t_screen_unchanged():
    seed()
    with H.browser() as b:
        pg = open_indent(b)
        rows = pg.eval_on_selector_all(
            "#indRows tr", "rs => rs.map(r => Array.from(r.cells).map(c => c.innerText.trim()))")
        first, second = rows[0], rows[1]
        assert first[1] == CUST_A and second[1] == "", (first, second)
        assert "item 2" in second[0], second[0]
        assert second[8].lower() == "open", second      # tags are upper-cased by CSS
        heads = pg.eval_on_selector_all("#v-indent thead th", "ts => ts.map(t => t.textContent.trim())")
        assert heads == ["Indent", "Customer", "Item", "Cell type", "Ordered", "KW",
                         "Dispatched", "Remaining", "Status", "Due", ""], heads


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")     # a failure message may hold a …
    width = max(len(n) for n, _ in _results)
    passed = failed = 0
    try:
        for name, fn in _results:
            try:
                fn()
                print("  PASS  %-*s" % (width, name))
                passed += 1
            except Exception as e:
                print("  FAIL  %-*s  %s" % (width, name, e))
                if "-v" in sys.argv:
                    traceback.print_exc()
                failed += 1
        print("\n%d passed, %d failed" % (passed, failed))
    finally:
        H.cleanup()
    sys.exit(1 if failed else 0)
