"""
ICON TRACE - custom (non-ICON) serial numbers, read from Excel.

An indent can carry the customer's OWN serial numbers instead of ICON's. Those
cannot be generated, so Planning loads them from a workbook. The workbook is
the layout Planning itself exports after an allocation (and that Unit-1's
BARCODE.py produced before it) - allocation_barcodes() in app.py:

    row 1   merged heading   "620W - 1440 NOS BOROSIL RENEWABLES LIMITED"
    row 2   S.NO. | BARCODE  repeated for each column pair
    row 3+  1000 rows per pair, then a new pair to the right

so a sheet exported from Planning, or one the customer fills in the same shape,
loads without being rearranged. Reading is by the BARCODE header, not by
position, so extra columns or a different number of pairs do no harm.

This module only PARSES and checks what a file can say about itself. It
touches no database: whether a serial is already in the master is the
server's to check (app.py), against the master.

Nothing is silently dropped or "fixed" beyond case and surrounding spaces:
every problem is named with its cell, and a file with any problem loads
nothing - an operator corrects the sheet and uploads it again.
"""

import io
import re

# What a serial may be. The system upper-cases every serial it is given
# (scans, look-ups, the barcode export), so a custom serial is stored in
# upper case; it is a single token of printable ASCII - what a scanner reads
# and a label prints. The length bounds are ASSUMPTIONS (DECISIONS.md, open):
# short enough to rule out a stray header or note, long enough for any real
# customer scheme seen so far.
MIN_LEN, MAX_LEN = 4, 40
MAX_SERIALS = 50000
MAX_PROBLEMS_SHOWN = 50

_TOKEN = re.compile(r"^[!-~]+$")                 # printable ASCII, no whitespace
_HEADER = re.compile(
    r"^(barcode|barcodes|serial|serial\s*no\.?|serial\s*number|"
    r"module\s*serial(\s*no\.?)?)$", re.I)
_ICON_SHAPED = re.compile(r"^ICON\d{3}[A-Z]\d{2}[0-9A-C]\d{7}$")


def _col(n):
    """0 -> A, 25 -> Z, 26 -> AA"""
    s = ""
    n += 1
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def normalise(v):
    """A cell's value as a serial candidate: (text, note). Excel hands a
    numeric serial back as a float - 100234.0 - which is the number typed, not
    a serial with a decimal point; leading zeros are already gone in that case
    and cannot be recovered (the sheet's column should be Text)."""
    if v is None:
        return None, None
    if isinstance(v, bool):
        return str(v).upper(), None
    if isinstance(v, float) and v == int(v):
        return str(int(v)), None
    if isinstance(v, (int, float)):
        return str(v), None
    return str(v).strip().upper(), None


def check(serial):
    """Why this text cannot be a custom serial, or None."""
    if not serial:
        return "is empty"
    if re.search(r"\s", serial):
        return "contains a space"
    if len(serial) < MIN_LEN:
        return "is shorter than %d characters" % MIN_LEN
    if len(serial) > MAX_LEN:
        return "is longer than %d characters" % MAX_LEN
    if not _TOKEN.match(serial):
        return "contains a character a barcode scanner cannot read"
    if _ICON_SHAPED.match(serial):
        return ("looks like an ICON serial number - ICON serials are generated "
                "by Planning on an indent that does not use custom serials")
    return None


def check_list(serials):
    """Problems for a list ALREADY normalised (the API re-checks what a
    client sends, never trusting the browser to have run parse()):
    [(index, serial, why)]. Duplicates are named against their first place."""
    problems, first = [], {}
    for i, s in enumerate(serials):
        why = check(s)
        if why:
            problems.append((i, s, why))
        elif s in first:
            problems.append((i, s, "appears twice (first at position %d)" % (first[s] + 1)))
        else:
            first[s] = i
    return problems


def parse(data):
    """Read a workbook. Returns
        {"ok": bool, "why": str (when not ok), "sheet", "heading",
         "serials": [...], "problems": [{"cell", "serial", "why"}],
         "problem_total": int}
    ok is False when no BARCODE column exists or anything is wrong."""
    try:
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as e:
        return {"ok": False, "why": "Could not read the workbook: %s" % e}

    found = None
    for ws in wb.worksheets:
        rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
        for ri in range(min(len(rows), 15)):
            cols = [ci for ci, v in enumerate(rows[ri])
                    if isinstance(v, str) and _HEADER.match(v.strip())]
            if cols:
                found = (ws.title, rows, ri, cols)
                break
        if found:
            break
    if not found:
        return {"ok": False, "why":
                "No BARCODE column found. Expected the layout Planning "
                "exports - a heading row, then S.NO. | BARCODE (repeated "
                "side by side for more than 1,000 serials)."}

    title, rows, hdr, cols = found
    heading = ""
    for r in rows[:hdr]:
        for v in r:
            if isinstance(v, str) and v.strip():
                heading = v.strip()
                break
        if heading:
            break

    serials, problems, seen = [], [], {}
    for ci in cols:
        for ri in range(hdr + 1, len(rows)):
            row = rows[ri]
            raw = row[ci] if ci < len(row) else None
            text, _ = normalise(raw)
            if text is None or text == "":
                continue
            cell = "%s%d" % (_col(ci), ri + 1)
            why = check(text)
            if why:
                problems.append({"cell": cell, "serial": text, "why": why})
            elif text in seen:
                problems.append({"cell": cell, "serial": text,
                                 "why": "appears twice (first at %s)" % seen[text]})
            else:
                seen[text] = cell
                serials.append(text)
            if len(serials) + len(problems) > MAX_SERIALS:
                return {"ok": False, "why": "More than %d serial numbers in one "
                        "file - split it." % MAX_SERIALS}
    if not serials and not problems:
        return {"ok": False, "why": "The BARCODE column is empty."}
    return {"ok": not problems, "sheet": title, "heading": heading,
            "serials": serials,
            "problems": problems[:MAX_PROBLEMS_SHOWN],
            "problem_total": len(problems)}
