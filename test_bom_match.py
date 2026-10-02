"""
ICON TRACE - the traceability file's bill of materials, read against the
material master.

    python test_bom_match.py

Mukesh compared a BOM backfilled from the monthly Excel with the same module's
BOM from Planning: the make read "LIONSOLAR 25.6% (210*182.2) G12R CELL" - make,
efficiency and size in one string, written as it stood, where the master says
"Lion Solar" - and most materials were "not recorded" because the importer read
only four of the file's columns.

THE RULES THIS FILE DEFENDS

  1. A make is resolved to the MASTER'S spelling; the efficiency, the size and
     the batch are taken OUT of the text. Typos (BORORSIL, YUEJIYA) resolve;
     what cannot be resolved is kept as written and reported, never dropped.
  2. Several makes / several batches are joined with ONE separator (",").
     A "/" inside a batch number (GG/26-27/422) is part of it.
  3. Every material the file names is read - each batch from the column beside
     it (both ribbons head theirs "Ribbon Invoice/Batch no") - against the form's
     own names (String Alignment Tape = Cell Alignment Tape, Channel & JB sealant
     = Sealant, EPE/POE = Encapsulant, the string ribbon fills Centre AND Edge,
     potting fills Part A AND B).
  4. Alternatives are chosen by the size the file states: Lead Bending Tape 15 or
     20 mm, Junction Box 0.4 or 0.3 mtr.
  5. A material the file is silent about, whose master has ONE make, takes it.
  6. What the master disagrees with is a NOTE: a size that differs, an unknown
     make; a cell efficiency the list lacks is ADDED to the list.
  7. A backfill batch's BOM can be refreshed from the file - never a Planning
     batch's, never one that covers a different set of serials.
"""

import io, os, sys, traceback, datetime

import ui_harness as H                                       # noqa: E402  (first: sets the DB path)
import store                                                 # noqa: E402
import db                                                    # noqa: E402
import auth_test_helper as AUTH                              # noqa: E402
import icon_bom_match as B                                   # noqa: E402
import icon_traceability_import as T                         # noqa: E402
import openpyxl                                              # noqa: E402

APP = H.APP
_results = []
os.environ.setdefault("ICON_PROD_BACKDATE_DAYS", "3650")

REAL = r"C:/Users/X/.claude/uploads/35caadef-138c-4dcc-b9c6-33c2f308233f/0ad28808-09___UNIT-2_TRACEABILITY_REPORT__SEPTEMBER_-2026_3.xlsx"

CELL = ["Lion Solar", "PT Nusa Solar", "Premier Energies", "Solar Space", "Tongwei Solar", "Yingfa"]
GLASS = ["Borosil", "Kibing", "Waaree", "Xinyi"]
FRAME = ["Aluvoltec", "Jiangsu Yuejia", "Jiangyin Yuanshuo (YS)", "Ralpro Techno", "Shanti Green", "Sudarshan"]
IC = ["Dhash", "Geba Copper", "Juren", "Sekhani Renewables", "Valeo"]


def test(name):
    def deco(fn):
        _results.append((name, fn))
        return fn
    return deco


# --------------------------------------------------------------------------
# 1,2  the matcher
# --------------------------------------------------------------------------

@test("makes resolve to the master's spelling - typos, extra words and sizes "
      "ignored - and several makes keep their order")
def t_makes():
    cases = [
        (CELL, "LIONSOLAR  25.7% (210*182.2) G12R CELL", ["Lion Solar"]),
        (CELL, "PREMIER 23.3%  (210*182.2) G12R CELL", ["Premier Energies"]),
        (GLASS, "BORORSIL (2376*1128*2 MM)", ["Borosil"]),
        (GLASS, "KIBING/ BORORSIL (2376*1128*2 MM)", ["Kibing", "Borosil"]),
        (IC, "GEBA  0.26MM", ["Geba Copper"]),
        (IC, "GEBA/DHASH  0.26MM", ["Geba Copper", "Dhash"]),
        (IC, "DHASH /SHEKHANI 0.26MM", ["Dhash", "Sekhani Renewables"]),
        (IC, "DHASH & GEBA 6.0X0.40", ["Dhash", "Geba Copper"]),
        (FRAME, "YS METAL  (2382*1134*30 MM)", ["Jiangyin Yuanshuo (YS)"]),
        (FRAME, "YUEJIYA (2382*1134*30 MM)", ["Jiangsu Yuejia"]),
        (FRAME, "ALUVOLTECH (2382*1134*30 MM)", ["Aluvoltec"]),
        (["H.B. Fuller", "Sunsol"], "HB FULLAR (15MM)", ["H.B. Fuller"]),
        (["Dhash", "GenX", "QC Solar"], "GNEX (0.4M) 30A ", ["GenX"]),
        (["Alishan", "Knack", "RenewSys", "Sheetsol"], "SHEETSOL EPE", ["Sheetsol"]),
    ]
    for makes, text, want in cases:
        got, left = B.match_makes(text, makes)
        assert got == want and not left, (text, got, left)


@test("a make the master does not have is kept as written and reported - never "
      "dropped, never forced onto a near name")
def t_unknown_make():
    got, left = B.match_makes("ZEBRA SOLAR (210*182.2) G12R CELL", CELL)
    assert got == [] and left == ["ZEBRA SOLAR"], (got, left)
    got, left = B.match_makes("LIONSOLAR/ ZEBRA", CELL)
    assert got == ["Lion Solar"] and left == ["ZEBRA"], (got, left)


@test("efficiency comes out of the text: single, both ends of a range, and "
      "never a size or a family name")
def t_efficiency():
    assert B.efficiencies("LIONSOLAR  25.7% (210*182.2) G12R CELL") == ["25.7%"]
    assert B.efficiencies("LIONSOLAR 25.6% -25.7%(210*182.2) G12R CELL") == ["25.6%", "25.7%"]
    assert B.efficiencies("LIONSOLAR  25.6-25.8% (210*182.2) G12R CELL") == ["25.6%", "25.8%"]
    assert B.efficiencies("PREMIER 23.3%  (210*182.2) G12R CELL") == ["23.3%"]
    assert B.efficiencies("LIONSOLAR (210*182.2) G12R CELL") == []


@test("batches use one separator whatever the file used, drop repeats, and "
      "keep a '/' that is part of a number")
def t_batches():
    assert B.clean_batch("ID20260709,ID20260709") == "ID20260709"
    assert B.clean_batch("ID20260721 & ID20260709") == "ID20260721,ID20260709"
    assert B.clean_batch("109255 & 109980") == "109255,109980"
    assert B.clean_batch("A / B") == "A,B"
    assert B.clean_batch("GG/26-27/422,GG/26-27/419") == "GG/26-27/422,GG/26-27/419"
    assert B.clean_batch("TIFAM2627/785") == "TIFAM2627/785"
    assert B.clean_batch(9000025061) == "9000025061" and B.clean_batch(9000025061.0) == "9000025061"
    assert B.clean_batch("  ") is None and B.clean_batch(None) is None


@test("sizes compare against the master's: same when the numbers agree whatever "
      "the order or punctuation, different when they do not, silent when the file "
      "states none")
def t_dims():
    assert B.same_dims("(210*182.2) G12R CELL", "182.2 x 210 mm \u00b7 >25% \u00b7 16 BB") is True
    assert B.same_dims("(2376*1128*2 MM)", "2376 x 1128 x 2 mm") is True
    assert B.same_dims("(2278*1134*30 MM)", "2382 x 1134 x 30 mm") is False
    assert B.same_dims("GNEX (0.4M) 30A", "0.4 mtr") is True
    assert B.same_dims("DHASH 4.0X0.42", "4.0 x 0.40 mm") is False
    assert B.same_dims("FASTO", "\u2014") is None


# --------------------------------------------------------------------------
# 3-6  the whole BOM, from a workbook in the real file's shape
# --------------------------------------------------------------------------

HEAD = ["Date", "Shift", "Module Wattage", "Serial No.", "Ending No", "Quantity",
        "Special Customer", "Remark", "Cell Make", "Cell Invoice/Batch no", "Cell Type",
        "Bus Bar", "Modules Sizes H & F", "Solar Glass (Front)",
        "Solar Glass (Front) Invoice/Batch no", "Solar Glass (Back)",
        "Solar Glass (Back) Invoice/Batch no", "Cells Inter connector",
        "Ribbon Invoice/Batch no", "String Inter Connector", "Ribbon Invoice/Batch no",
        "Flux", "Flux Invoice/Batch no", "EPE/POE", "EPE/POE Invoice/Batch no",
        "String Alignment Tape", "String Alignment Tape Invoice/Batch no",
        "Lead Bending Tape", "Lead Bending Tape Invoice/Batch no",
        "Edge Sealing Tape", "Edge Sealing Tape Invoice/Batch no", "Junction Box",
        "Junction Box Invoice/Batch no", "Channel & JB sealant", "Sealant Invoice/Batch no",
        "Potting Material", "Potting Invoice/Batch no", "Aluminium Frame",
        "Frame Invoice/Batch no", "RFID", "RFID Invoice/Batch no", "Shift Incharge", "Remark"]

DEFAULT = {"cell": "LIONSOLAR  25.7% (210*182.2) G12R CELL", "cell_b": "ID20260701,ID20260701",
           "type": "G12R", "bb": "16 BB", "size": "790, 1400, 1094 MM",
           "gf": "KIBING/ BORORSIL (2376*1128*2 MM)", "gf_b": "KBM-260408 & 9000025061",
           "gb": "BORORSIL (2376*1128*2 MM)", "gb_b": 9000025061,
           "cic": "GEBA/DHASH  0.26MM", "cic_b": 113166,
           "sic": "GEBA 6.0X0.40 AND  DHASH 4.0X0.42", "sic_b": "109255 & 109980",
           "flux": "RCPV", "flux_b": "FY26-27/119",
           "epe": "ALISHAN/SHEETSOL", "epe_b": "6697/26-27/754",
           "at": "HB FULLAR (8MM)", "at_b": "2026-27/17",
           "lt": "HB FULLAR (15MM)", "lt_b": "2026-27/17",
           "et": "HB FULLAR (30MM)", "et_b": "2026-27/16",
           "jb": "GNEX (0.4M) 30A ", "jb_b": "GX03092627310",
           "sl": "FASTO", "sl_b": "TIFAM2627/785", "pt": "FASTO", "pt_b": "TIFAM2627/785",
           "fr": "YS METAL  (2382*1134*30 MM)", "fr_b": "YJMI-IC2601",
           "rf": "FINOTECH", "rf_b": "FY26-27/118"}
ORDER = ["cell", "cell_b", "type", "bb", "size", "gf", "gf_b", "gb", "gb_b", "cic", "cic_b",
         "sic", "sic_b", "flux", "flux_b", "epe", "epe_b", "at", "at_b", "lt", "lt_b",
         "et", "et_b", "jb", "jb_b", "sl", "sl_b", "pt", "pt_b", "fr", "fr_b", "rf", "rf_b"]


def S(i):
    return "ICON625R1290140%03d" % i


def workbook(rows=None, **over):
    """The real file's shape: empty column A, a title block, 43 headers, one row."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TRACEABILITY TEST"
    ws.cell(3, 7, "Shiftwise ... for the month TEST - 2026")
    for c, h in enumerate(HEAD):
        ws.cell(12, 2 + c, h)
    rows = rows or [dict(over)]
    for n, ov in enumerate(rows):
        v = dict(DEFAULT)
        v.update(ov)
        vals = [datetime.datetime(2026, 9, 1), "A", "625W", v.get("start", S(1)), v.get("end", S(7)),
                v.get("qty", 7), "BOROSIL", True] + [v[k] for k in ORDER] + ["YAMAN & RAJKUMAR", None]
        for c, x in enumerate(vals):
            if x is not None:
                ws.cell(13 + n, 2 + c, x)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def catalog():
    with store.conn() as (cx, cur):
        db.seed_materials(cur)
        return db.materials(cur), db.cell_efficiencies(cur)


def bom_of(data, **kw):
    r = T.parse(data)
    assert r["ok"] and not r["problems"], r
    cat, known = catalog()
    return T.bom_materials(r["ranges"][0]["bom"], cat, wattage=625, known_efficiencies=known), cat, r


def by_no(out):
    return {m["material_no"]: m for m in out["rows"]}


@test("every material the file names is read, each batch from the column beside "
      "it - including the two ribbons that share a header - under the form's own names")
def t_full_bom():
    store.wipe()
    out, cat, r = bom_of(workbook())
    items = r["ranges"][0]["bom"]["items"]
    assert set(items) == {"cell", "glass_front", "glass_back", "cell_ic", "string_ic", "flux",
                          "epe", "align_tape", "lead_tape", "edge_tape", "jb", "sealant",
                          "potting", "frame", "rfid"}, sorted(items)
    assert items["cell_ic"]["batch"] == "113166" and items["string_ic"]["batch"] == "109255 & 109980"
    b = by_no(out)
    assert (b[5]["vendor"], b[5]["efficiency"], b[5]["batch"]) == ("Lion Solar", "25.7%", "ID20260701"), b[5]
    assert (b[6]["vendor"], b[6]["batch"]) == ("Kibing,Borosil", "KBM-260408,9000025061"), b[6]
    assert (b[7]["vendor"], b[7]["batch"]) == ("Borosil", "9000025061"), b[7]
    assert (b[12]["vendor"], b[12]["batch"]) == ("Geba Copper,Dhash", "113166"), b[12]
    # one ribbon column fills Centre AND Edge: "<centre> AND <edge>", batches shared
    assert b[13]["vendor"] == "Geba Copper" and b[14]["vendor"] == "Dhash", (b[13], b[14])
    assert b[13]["batch"] == b[14]["batch"] == "109255,109980"
    assert b[9]["vendor"] == "Alishan,Sheetsol" and b[19]["vendor"] == "RCPV"
    assert b[20]["vendor"] == "H.B. Fuller" and b[22]["vendor"] == "H.B. Fuller"
    assert b[16]["vendor"] == "Fasto" and b[17]["vendor"] == b[18]["vendor"] == "Fasto"
    assert b[8]["vendor"] == "Jiangyin Yuanshuo (YS)" and b[15]["vendor"] == "Finotech"
    assert b[10]["vendor"] == "GenX" and b[10]["batch"] == "GX03092627310"


@test("alternatives follow the size the file states: tape 15 -> the 15 mm tape, 20 "
      "-> the 20 mm one; junction box 0.4 / 0.3 mtr")
def t_alternatives():
    store.wipe()
    cat, _ = catalog()
    n15 = [m["n"] for m in cat if m["name"] == "Lead Bending Tape" and "15" in m["size"]]
    n20 = [m["n"] for m in cat if m["name"] == "Lead Bending Tape" and "20" in m["size"]]
    assert n15 and n20, "the catalog has no 15 mm / 20 mm lead bending tape"
    assert all(m.get("group") == "LBT" for m in cat if m["name"] == "Lead Bending Tape"), "not alternatives"
    for text, want in (("HB FULLAR (15MM)", n15[0]), ("HB FULLAR (20MM)", n20[0])):
        out, _, _ = bom_of(workbook(lt=text))
        nums = [m["material_no"] for m in out["rows"] if m["material_no"] in (n15[0], n20[0])]
        assert nums == [want], (text, nums)
    out, _, _ = bom_of(workbook(jb="GNEX (0.3M) 30A"))
    assert 11 in by_no(out) and 10 not in by_no(out), "0.3 mtr junction box not chosen"


@test("a material the file is silent about takes its single make (EPE Strip, the "
      "625 W back label) - and one with several makes is left unrecorded")
def t_single_make_default():
    store.wipe()
    out, cat, _ = bom_of(workbook())
    b = by_no(out)
    assert b[23]["vendor"] == "RenewSys" and b[28]["vendor"] == "Kvell", sorted(b)
    assert 24 not in b and 27 not in b, "another wattage's back label was recorded"
    for n in (29, 30):          # barcode label, pallet packing: several makes
        assert n not in b, n
    assert any(x["kind"] == "default" and x["material"] == "EPE Strip (Output Patti)" for x in out["notes"])


@test("what the master disagrees with is a NOTE and the make is still recorded: "
      "a differing size, an unknown make; new efficiencies are reported")
def t_notes():
    store.wipe()
    out, _, _ = bom_of(workbook(fr="ALUVOLTECH (2278*1134*30 MM)", sic="DHASH 6.0X0.40 AND DHASH 4.0X0.42",
                                 cell="PREMIER 23.3% (210*182.2) G12R CELL",
                                 gb="ZEBRA GLASS (2376*1128*2 MM)"))
    b = by_no(out)
    kinds = {(n["kind"], n["material"]) for n in out["notes"]}
    assert ("size", "Aluminium Frame") in kinds and ("size", "String Inter Connector \u2014 Edge") in kinds, kinds
    assert ("make", "Solar Glass \u2014 Rear") in kinds, kinds
    assert b[8]["vendor"] == "Aluvoltec" and b[7]["vendor"] == "ZEBRA GLASS", (b[8], b[7])
    assert out["new_efficiencies"] == ["23.3%"], out["new_efficiencies"]
    assert b[5]["efficiency"] == "23.3%" and b[5]["vendor"] == "Premier Energies"


# --------------------------------------------------------------------------
# 6,7  applied
# --------------------------------------------------------------------------

def client():
    store.wipe()
    c = APP.app.test_client()
    AUTH.test_login(c)
    return c


def parse_up(c, data):
    return c.post("/api/prodentry/import/parse", data={"file": (io.BytesIO(data), "t.xlsx")},
                  content_type="multipart/form-data")


def apply(c, rng, backfill=True):
    return c.post("/api/prodentry/import/apply",
                  json={"ranges": [rng], "incharge": "NIGHT INCHARGE", "backfill": backfill,
                        "dcr": "NDCR"}).get_json()


def rows_of(aid):
    with store.conn() as (cx, cur):
        return {r["material_no"]: dict(r) for r in store.rows(
            cur, "SELECT * FROM allocation_material WHERE alloc_id=%s", (aid,))}


@test("backfill writes the complete, matched BOM, and ADDS a cell efficiency the "
      "master's list lacks - in numeric order - and says so")
def t_apply_writes_matched_bom():
    c = client()
    data = workbook(cell="LIONSOLAR 25.8% (210*182.2) G12R CELL")
    p = parse_up(c, data).get_json()
    assert p["bom_summary"]["new_efficiencies"] == ["25.8%"], p["bom_summary"]
    r = apply(c, p["ranges"][0])
    assert r["recorded"] == 1 and r["efficiencies_added"] == ["25.8%"], r
    b = rows_of(r["results"][0]["alloc_id"])
    assert len(b) >= 19 and b[5]["vendor"] == "Lion Solar" and b[5]["efficiency"] == "25.8%", sorted(b)
    with store.conn() as (cx, cur):
        eff = db.cell_efficiencies(cur)
    assert "25.8%" in eff and eff == sorted(eff, key=lambda v: float(v.rstrip("%"))), eff
    assert eff.index("25.8%") == len(eff) - 1


@test("the parse step reports, before anything is applied, what the master "
      "disagrees with across the whole file")
def t_parse_summary():
    c = client()
    data = workbook(rows=[{}, {"start": S(8), "end": S(14), "fr": "ALUVOLTECH (2278*1134*30 MM)"}])
    s = parse_up(c, data).get_json()["bom_summary"]
    notes = {(n["kind"], n["material"]): n["ranges"] for n in s["notes"]}
    assert notes[("size", "Aluminium Frame")] == 1 and notes[("size", "String Inter Connector \u2014 Edge")] == 2, notes
    assert not any(k[0] == "default" for k in notes)


@test("re-importing a backfilled range REFRESHES its BOM from the file - the repair "
      "for BOMs recorded as raw text - and an unchanged file changes nothing")
def t_refresh_backfill():
    c = client()
    p = parse_up(c, workbook()).get_json()
    aid = apply(c, p["ranges"][0])["results"][0]["alloc_id"]
    good = rows_of(aid)
    with store.conn() as (cx, cur):          # as the old importer left it
        cur.execute("DELETE FROM allocation_material WHERE alloc_id=%s AND material_no<>5", (aid,))
        cur.execute("UPDATE allocation_material SET vendor='LIONSOLAR 25.6% (210*182.2) G12R CELL' "
                    "WHERE alloc_id=%s AND material_no=5", (aid,))
    r = apply(c, p["ranges"][0])
    res = r["results"][0]
    assert res["action"] == "bom_refreshed" and r["bom_refreshed"] == 1 and r["recorded"] == 0, r
    assert rows_of(aid) == good, "the refreshed BOM is not the matched one"
    again = apply(c, p["ranges"][0])["results"][0]
    assert again["action"] == "bom_refreshed" and rows_of(aid) == good
    with store.conn() as (cx, cur):
        assert store.one(cur, "SELECT COUNT(*) AS n FROM production_entry")["n"] == 1, "a refresh wrote production"
        assert store.one(cur, "SELECT COUNT(*) AS n FROM serial")["n"] == 7


@test("a Planning batch's BOM is never touched, and nor is a batch covering a "
      "different set of serials")
def t_refresh_refused():
    c = client()
    # (a) serials issued by Planning
    with store.conn() as (cx, cur):
        iid = store.insert(cur, "indent", {"indent_no": "P/1", "indent_date": "2026-09-01",
                                           "customer": "C0002", "created_by": "t"})
        lid = store.insert(cur, "indent_line", {"indent_id": iid, "line_no": 1, "model": "ISEN625-G12R",
            "item_description": "x", "wattage": 625, "qty": 7, "dcr": "NDCR"})
        aid = store.insert(cur, "allocation", {"indent_line_id": lid, "model": "ISEN625-G12R",
            "wattage": 625, "customer": "C0002", "date_produced": "2026-09-01", "shift": 1, "qty": 7,
            "seq_from": 1, "seq_to": 7, "created_by": "t"})
        store.insert(cur, "allocation_material", {"alloc_id": aid, "material_no": 5, "vendor": "Premier Energies"})
        for i in range(1, 8):
            store.insert(cur, "serial", {"serial": S(i), "build_instance": 1, "alloc_id": aid,
                "indent_line_id": lid, "model": "ISEN625-G12R", "wattage": 625, "customer": "C0002",
                "dcr": "NDCR", "format_version": 2, "date_produced": "2026-09-01", "shift": 1,
                "sequence": i, "state": "planned"})
    p = parse_up(c, workbook()).get_json()
    res = apply(c, p["ranges"][0])["results"][0]
    assert res["action"] == "skipped" and "made in Planning" in res["why"], res
    assert list(rows_of(aid)) == [5] and rows_of(aid)[5]["vendor"] == "Premier Energies"
    # (b) a backfill batch that covers MORE than this range
    c = client()
    big = parse_up(c, workbook(start=S(1), end=S(7))).get_json()["ranges"][0]
    aid = apply(c, big)["results"][0]["alloc_id"]
    sub = dict(big, start=S(1), end=S(3), qty=3, seq_to=3)
    res = apply(c, sub)["results"][0]
    assert res["action"] == "skipped" and "covers 7 serials" in res["why"], res


# --------------------------------------------------------------------------
# the real file
# --------------------------------------------------------------------------

@test("the real September file: every make resolves to the master, nothing is "
      "missing, and the batch Mukesh compared reads like Planning's own")
def t_real_file():
    if not os.path.exists(REAL):
        print("      (real file not on this machine - skipped)")
        return
    store.wipe()
    cat, known = catalog()
    r = T.parse(open(REAL, "rb").read())
    assert r["ok"] and r["summary"]["ranges"] == 275, r["summary"]
    kinds, neweff = {}, set()
    for rg in r["ranges"]:
        out = T.bom_materials(rg["bom"], cat, wattage=rg["wattage"], known_efficiencies=known)
        neweff.update(out["new_efficiencies"])
        for n in out["notes"]:
            kinds[n["kind"]] = kinds.get(n["kind"], 0) + 1
            assert n["kind"] in ("default", "size"), n      # no unknown makes, no missing materials
        assert len(out["rows"]) >= 18, (rg["start"], len(out["rows"]))
    assert neweff == {"23.3%", "25.8%"}, neweff
    rg = [x for x in r["ranges"] if x["start"] == "ICON625R1293030001"][0]
    b = by_no(T.bom_materials(rg["bom"], cat, wattage=625, known_efficiencies=known))
    assert (b[5]["vendor"], b[5]["efficiency"]) == ("Lion Solar", "25.6%"), b[5]
    assert b[6]["vendor"] == "Borosil" and b[8]["vendor"] == "Jiangyin Yuanshuo (YS)"
    assert 31 in b and 21 not in b, "the 15 mm lead bending tape was not chosen"


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    only = [a for a in sys.argv[1:] if not a.startswith("-")]
    width = max(len(n) for n, _ in _results)
    passed = failed = 0
    try:
        for name, fn in _results:
            if only and not any(o in fn.__name__ for o in only):
                continue
            try:
                fn()
                print("  PASS  %-*s" % (width, name))
                passed += 1
            except Exception as e:
                print("  FAIL  %-*s  %s" % (width, name, str(e)[:400]))
                if "-v" in sys.argv:
                    traceback.print_exc()
                failed += 1
        print("\n%d passed, %d failed" % (passed, failed))
    finally:
        H.cleanup()
    sys.exit(1 if failed else 0)
