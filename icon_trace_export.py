"""
ICON TRACE - the monthly traceability report, as an Excel export.

The layout is the plant's own report (UNIT-2 TRACEABILITY REPORT, document
IS-MP2-PDN-FM-02): a title block, a production summary by wattage, then one row
per production run - date, shift, wattage, serial range, quantity, customer -
and what the module was built from.

ONE MATERIAL, ONE COLUMN. The hand-kept report folds two materials into one
column (both potting parts, the centre and the edge ribbon, the cell's make and
its efficiency). This export never does: every material of the master has its own
Make column and its own Invoice/Batch column, and a cell's efficiency is its own
column too. A group of alternatives (Lead Bending Tape 20 / 15 mm, the edge ribbon
0.40 / 0.41 / 0.42, the frame's mounting holes) shares ONE column - a module
carries one of them - and the variant that was used is stated beside the make.

Pure: rows and the material master in, workbook bytes out. The route in app.py
does the reading. Nothing here touches the database.
"""

import datetime
import io
import re

JOINER = ","                  # the system-wide separator for several makes / batches
DOC_NO = "Document No:IS-MP2-PDN-FM-02"
REVIEW = "Review Date & No:  01.03.2026 & 00"
TITLE = "ICON SOLAR-EN POWER TECHNOLOGIES PVT. LTD.  -  (TRACEABILITY)"

# the wattages the summary lists when nothing else was built, as the report does
STANDARD_WATTS = (625, 620, 630, 590)
WATT_LINES = 4               # rows 6..9 hold four wattage lines under the total in row 5
FIRST_ROW = 13               # the report's first data row (headers on 12)

_MONTHS = ["JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
           "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER"]
_LABEL = re.compile(r"\s+\d+\s*WP$", re.I)
_BB = re.compile(r"(\d+)\s*BB", re.I)

# What a spreadsheet treats as the start of a formula. A batch, a make or a name
# read from an imported file may begin with one ("=HYPERLINK(...)"): it is written
# as TEXT, never evaluated when somebody opens the report.
_FORMULA_LEAD = ("=", "+", "-", "@", "\t", "\r")

# the report's own check on a row: the quantity matches the serial span
CHECK = ('=IF($G{r}="","",IF(OR($G{r}=1,IF(ISERROR(RIGHT($F{r},4)-RIGHT($E{r},4)),'
         'FALSE,(RIGHT($F{r},4)-RIGHT($E{r},4)+1)=$G{r})),TRUE,FALSE))')


def column_name(m):
    """The column a master material belongs to: its name, with a back label's
    wattage dropped (the wattage has its own column)."""
    if (m.get("series") or "") == "LABEL":
        return _LABEL.sub("", m["name"]).strip()
    return m["name"]


def _head(size):
    """A master size without what follows its first '·' (the mounting holes,
    '3 hole · grid printed') - '2382 x 1134 x 30 mm'."""
    h = str(size or "").split("·")[0].strip()
    return "" if h in ("—", "-") else h


def holes_of(size):
    """'790, 1400, 1094 MM' from a master size that carries its holes."""
    m = re.search(r"holes?\s*([\d.,\s]+)", str(size or ""), re.I)
    if not m:
        return None
    nums = re.findall(r"\d+(?:\.\d+)?", m.group(1))
    return ", ".join(nums) + " MM" if nums else None


def material_groups(catalog, categories, recorded=()):
    """[{name, members: [n..], eff, ...}] in the BOM screen's order - category,
    then master number. A group made only of legacy materials is listed only when
    something in the export was recorded against it: a recorded fact is never
    left off the sheet."""
    rank = {c: i for i, c in enumerate(categories)}
    order = sorted(catalog, key=lambda m: (rank.get(m.get("cat"), 99), m["n"]))
    groups, by_name = [], {}
    for m in order:
        name = column_name(m)
        g = by_name.get(name)
        if g is None:
            g = by_name[name] = {"name": name, "members": [], "eff": False,
                                 "live": False, "sized": []}
            groups.append(g)
        g["members"].append(m["n"])
        g["eff"] = g["eff"] or bool(m.get("cell"))
        g["live"] = g["live"] or not m.get("legacy")
        g["sized"].append(m)
    recorded = set(recorded)
    return [g for g in groups if g["live"] or recorded & set(g["members"])]


def columns(groups):
    """The sheet's columns, in order: [(key, header, width)]. The report's lead
    columns, each material's Make (and Efficiency) and Invoice/Batch columns, the
    three descriptive columns after the cell, then the incharge and remark."""
    cols = [("date", "Date", 11.5), ("shift", "Shift", 8), ("watt", "Module Wattage", 12.7),
            ("start", "Serial No.", 25.7), ("end", "Ending No", 25.7),
            ("qty", "Quantity", 9.5), ("customer", "Special Customer", 30),
            ("check", "Remark", 10)]
    for g in groups:
        n = g["name"]
        cell = g["eff"]
        cols.append(("mk:" + n, "Cell Make" if cell else n, 30 if not cell else 28))
        if cell:
            cols.append(("ef:" + n, "Cell Efficiency", 12))
        cols.append(("bt:" + n, "Cell Invoice/Batch no" if cell else n + " Invoice/Batch no", 26))
        if cell:
            cols += [("celltype", "Cell Type", 10), ("bb", "Bus Bar", 9),
                     ("holes", "Modules Sizes H & F", 24)]
    cols += [("incharge", "Shift Incharge", 34), ("remark", "Remark", 14)]
    return cols


def _model_family(models, model):
    m = models.get(model) or {}
    return m.get("family") or (model.rsplit("-", 1)[-1] if "-" in (model or "") else "")


def cells_of(row, groups, by_no, models):
    """{column key: value} for one report row."""
    out = {"date": row["date"], "shift": row["shift"],
           "watt": "%dW" % int(row["wattage"]), "start": row["start"], "end": row["end"],
           "qty": row["qty"], "customer": row["customer"],
           "incharge": row.get("incharge") or "",
           "remark": "REWORK" if row.get("rework") else ""}
    bom = row.get("bom") or {}
    out["celltype"] = _model_family(models, row["model"])
    frame_holes = bus = None
    for g in groups:
        got = [(n, bom[n]) for n in g["members"] if n in bom]
        if not got:
            continue
        make = JOINER.join(v["vendor"] for _n, v in got if v.get("vendor"))
        batch = JOINER.join(v["batch"] for _n, v in got if v.get("batch"))
        if g["eff"]:
            out["ef:" + g["name"]] = JOINER.join(v["efficiency"] for _n, v in got if v.get("efficiency"))
            m = by_no.get(got[0][0]) or {}
            hit = _BB.search(str(m.get("size") or ""))
            bus = ("%s BB" % hit.group(1)) if hit else bus
        elif len(g["members"]) > 1 and (g["sized"][0].get("series") or "") != "LABEL":
            sz = _head(by_no.get(got[0][0], {}).get("size"))
            if make and sz:
                make = "%s · %s" % (make, sz)
        if g["name"] == "Aluminium Frame":
            frame_holes = holes_of(by_no.get(got[0][0], {}).get("size")) or frame_holes
        out["mk:" + g["name"]] = make
        out["bt:" + g["name"]] = batch
    out["bb"] = bus or ""
    # the frame's own mounting holes where it states them; the model master's size otherwise
    out["holes"] = frame_holes or (models.get(row["model"]) or {}).get("size") or ""
    return out


def summary_watts(rows):
    """(wattages with a line of their own, whether the rest go to 'other'). Those
    built, most first; the report's usual ones fill what is left of the four lines
    (so a month with only 625 W still shows 620, 630, 590 at 0, as the report does)."""
    qty = {}
    for r in rows:
        qty[int(r["wattage"])] = qty.get(int(r["wattage"]), 0) + int(r["qty"] or 0)
    built = sorted(qty, key=lambda w: (-qty[w], w))
    if len(built) > WATT_LINES:
        return built[:WATT_LINES - 1], True
    pad = [w for w in STANDARD_WATTS if w not in built]
    return built + pad[:WATT_LINES - len(built)], False


def period(dfrom, dto):
    """(sheet label, subtitle, kind) - kind is 'month', 'day' or 'range'."""
    if dfrom == dto:
        return ("TRACEABILITY %s" % dfrom.strftime("%d.%m.%Y"),
                "Shiftwise cell wattage and serial Numbers for %s" % dfrom.strftime("%d-%m-%Y"), "day")
    last = (dfrom.replace(day=28) + datetime.timedelta(days=4)).replace(day=1) - datetime.timedelta(days=1)
    if dfrom.day == 1 and dto == last:
        return ("TRACEABILITY %s-%d" % (_MONTHS[dfrom.month - 1][:3], dfrom.year),
                "Shiftwise cell wattage and serial Numbers for the month %s - %d"
                % (_MONTHS[dfrom.month - 1], dfrom.year), "month")
    return ("TRACEABILITY %s-%s" % (dfrom.strftime("%d.%m.%y"), dto.strftime("%d.%m.%y")),
            "Shiftwise cell wattage and serial Numbers from %s to %s"
            % (dfrom.strftime("%d-%m-%Y"), dto.strftime("%d-%m-%Y")), "range")


def filename(dfrom, dto):
    if dfrom == dto:
        return "UNIT-2 TRACEABILITY REPORT %s.xlsx" % dfrom.strftime("%d-%m-%Y")
    _l, _s, kind = period(dfrom, dto)
    if kind == "month":
        return "UNIT-2 TRACEABILITY REPORT %s-%d.xlsx" % (_MONTHS[dfrom.month - 1], dfrom.year)
    return "UNIT-2 TRACEABILITY REPORT %s to %s.xlsx" % (dfrom.strftime("%d-%m-%Y"), dto.strftime("%d-%m-%Y"))


def build(rows, catalog, categories, models, dfrom, dto):
    """The workbook, as bytes. `rows`: one dict per report row - date (a date),
    shift, wattage, model, start, end, qty, customer, incharge, rework, bom
    ({material_no: {vendor, efficiency, batch}})."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L

    by_no = {m["n"]: m for m in catalog}
    recorded = {n for r in rows for n in (r.get("bom") or {})}
    groups = material_groups(catalog, categories, recorded)
    cols = columns(groups)
    label, subtitle, kind = period(dfrom, dto)

    wb = Workbook()
    ws = wb.active
    ws.title = label[:31]
    ws.sheet_view.showGridLines = False
    ws.sheet_view.zoomScale = 90
    navy, pale, zebra = "1F4E79", "E9EDF4", "F7F9FB"
    thin, med = Side(style="thin"), Side(style="medium")
    box = Border(left=thin, right=thin, top=thin, bottom=thin)
    mid = Alignment(horizontal="center", vertical="center", wrap_text=True)

    first, last_col = 2, 1 + len(cols)            # column A stays empty, as in the report
    for i, (_k, _h, w) in enumerate(cols):
        ws.column_dimensions[L(first + i)].width = w
    ws.column_dimensions["A"].width = 2

    # ---- title block
    right = min(last_col, first + 12)
    ws.merge_cells(start_row=2, start_column=7, end_row=2, end_column=right)
    t = ws.cell(2, 7, TITLE)
    t.font = Font(name="Calibri", size=16, bold=True, color="FFFFFF")
    t.fill = PatternFill("solid", fgColor=navy)
    t.alignment = mid
    ws.merge_cells(start_row=3, start_column=7, end_row=3, end_column=right)
    s = ws.cell(3, 7, subtitle)
    s.font = Font(name="Calibri", size=14)
    s.fill = PatternFill("solid", fgColor=pale)
    s.alignment = mid
    ws.row_dimensions[1].height = 9
    ws.row_dimensions[2].height = 21
    ws.row_dimensions[3].height = 18
    ws.row_dimensions[11].height = 1.5
    ws.row_dimensions[12].height = 32.4

    # ---- summary: the total, then the wattages
    n_rows = len(rows)
    lastr = FIRST_ROW + max(n_rows, 1) - 1
    qcol = L(first + [k for k, _h, _w in cols].index("qty"))
    wcol = L(first + [k for k, _h, _w in cols].index("watt"))
    shown, extra = summary_watts(rows)
    ws.merge_cells("B4:F4")
    h = ws["B4"]
    h.value = "Monthly Traceability Summary" if kind == "month" else "Traceability Summary"
    h.font = Font(name="Calibri", size=12, bold=True, color=navy)
    h.alignment = Alignment(horizontal="left", vertical="center")

    def line(r, text, formula):
        ws.merge_cells(start_row=r, start_column=2, end_row=r, end_column=4)
        a = ws.cell(r, 2, text)
        a.font = Font(name="Calibri", size=12)
        a.alignment = Alignment(horizontal="left", vertical="center")
        v = ws.cell(r, 5, formula)
        v.font = Font(name="Calibri", size=12, bold=True)
        v.alignment = Alignment(horizontal="right", vertical="center")
        v.number_format = "0"
        u = ws.cell(r, 6, "Modules")
        u.font = Font(name="Calibri", size=12)
        for c in range(2, 7):
            ws.cell(r, c).fill = PatternFill("solid", fgColor=pale)
            ws.cell(r, c).border = Border(top=thin, bottom=thin)

    rng_q = "%s%d:%s%d" % (qcol, FIRST_ROW, qcol, lastr)
    rng_w = "%s%d:%s%d" % (wcol, FIRST_ROW, wcol, lastr)
    line(5, "Total Production (Modules)" if kind != "month" else "Total Monthly Production (Modules)",
         "=SUM(%s)" % rng_q)
    r = 6
    for w in shown:
        line(r, "Total %dW Production" % w, '=SUMIF(%s,"%dW",%s)' % (rng_w, w, rng_q))
        r += 1
    if extra:                  # more wattages than the report's block has lines for
        line(r, "Total other wattages", "=E5-" + "-".join("E%d" % k for k in range(6, r)))
    ws.merge_cells(start_row=6, start_column=last_col - 2, end_row=7, end_column=last_col)
    ws.cell(6, last_col - 2, DOC_NO).alignment = mid
    ws.merge_cells(start_row=9, start_column=last_col - 2, end_row=10, end_column=last_col)
    ws.cell(9, last_col - 2, REVIEW).alignment = mid
    for rr in (6, 9):
        ws.cell(rr, last_col - 2).font = Font(name="Calibri", size=13)

    # ---- headers
    for i, (_k, hd, _w) in enumerate(cols):
        c = ws.cell(12, first + i, hd)
        c.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=navy)
        c.alignment = mid
        c.border = box

    # ---- rows
    keys = [k for k, _h, _w in cols]
    for n, row in enumerate(rows):
        r = FIRST_ROW + n
        cells = cells_of(row, groups, by_no, models)
        for i, k in enumerate(keys):
            if k == "check":
                v = CHECK.format(r=r)
            else:
                v = cells.get(k, "")
            c = ws.cell(r, first + i, v if v != "" else None)
            if k != "check" and isinstance(v, str) and v[:1] in _FORMULA_LEAD:
                c.data_type = "s"          # text that merely STARTS like a formula stays text
            c.font = Font(name="Calibri", size=11)
            c.alignment = mid
            c.border = box
            if k == "date":
                c.number_format = "dd-mm-yyyy"
            if n % 2:
                c.fill = PatternFill("solid", fgColor=zebra)
        ws.row_dimensions[r].height = 16.2
    ws.freeze_panes = "C%d" % FIRST_ROW
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    from openpyxl.worksheet.properties import PageSetupProperties
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_title_rows = "12:12"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
