"""
ICON TRACE - tests for EL Analysis (icon_el_analysis.py, /api/el/analysis).

    python test_el_analysis.py

THE RULES THIS FILE DEFENDS (the EL Insight Roadmap and the EL Defect Master,
approved 17 Oct 2026; the correctness rework Mukesh approved on 19 Oct 2026)

  1. A window is what its button says - the last hour in 5-minute bars, 8 hours
     and a shift by the hour, a factory day (06:00 to 06:00) by the hour, 7 days
     by day and shift, 30 days by day - and its change is read against the
     window of the same length just before it, when the index covers at least
     95% of that window.

  2. An image counts by when it was taken. One a machine in continuous mode
     passed by itself is counted apart and is in no NG rate. A re-filed image is
     one image, its latest filing standing.

  3. A module's first image at a station is its FIRST image; every later one a
     RE-IMAGE. The NG rate, every split and every comparison are on first
     images; re-images have their own figure ("NG again").

  4. A comparison is of RATES, with its significance: a busier window is not a
     rise. Control limits on the trend are Laney p'.

  5. A reason is read through the one defect list - every spelling one reason; a
     folder the list does not have is a reason of its own, marked new. The
     Pareto ranks modules (NG at their first image), Other in its ranked place.

  6. Modules beside images; stations never added together. FQC's decision is
     shown by the EL reason it read.

  7. Every figure opens its modules.

  8. Production Incharge, Admin and Super Admin see it; others by permission.

Each test names the rule it defends, so a failure says which decision broke.
"""

import datetime, io, os, shutil, sys, tempfile, traceback

TMP = tempfile.mkdtemp(prefix="icontrace_elan_")
os.environ["ICON_DB_FILE"] = os.path.join(TMP, "test.db")

import db                                                    # noqa: E402
import store                                                 # noqa: E402
import icon_el_analysis as ela                               # noqa: E402
import app as APP                                            # noqa: E402
import auth_test_helper as AUTH                              # noqa: E402

_results = []


def test(name):
    def deco(fn):
        _results.append((name, fn))
        return fn
    return deco


NOW = datetime.datetime(2026, 10, 17, 15, 30, 0)        # B shift of the 17th
CFG = {"el_root": "", "el_a_root": r"X:\final_a", "el_b_root": r"X:\final_b",
       "pre_el_a_root": "", "pre_el_b_root": r"X:\pre_b1", "pre_el2_a_root": "",
       "pre_el2_b_root": r"X:\pre_b2", "post_el_a_root": "", "post_el_b_root": r"X:\post_b"}
_n = [0]


def fresh():
    store.wipe()
    _n[0] = 0
    with store.conn() as (cx, cur):
        db.set_config(cur, CFG)


def _row(cur, station, serial, verdict, at, unit=1, line="B", mode=None, recapture=0, seen=None):
    _n[0] += 1
    cur.execute("INSERT INTO el_capture (station, unit, line, serial, recapture, verdict, path, "
                "file_time, seen_at, mode) VALUES (%s,%s,%s,%s,%s,%s,%s,%s,%s,%s)",
                (station, unit, line, serial, recapture, verdict,
                 r"X:\%s\%d\%s.jpg" % (station, _n[0], serial), at,
                 seen or "2026-10-17T15:%02d:%02d" % (_n[0] // 60 % 60, _n[0] % 60), mode))


def cap(station, serial, verdict, at, **kw):
    """One el_capture row, as the poller writes it."""
    with store.conn() as (cx, cur):
        _row(cur, station, serial, verdict, at, **kw)


def caps(rows):
    """Many rows in one go: (station, serial, verdict, at)."""
    with store.conn() as (cx, cur):
        for r in rows:
            _row(cur, *r)


def master(serial, model="ISEN625-G12R", customer="STOCK"):
    with store.conn() as (cx, cur):
        store.insert(cur, "serial", {"serial": serial, "build_instance": 1, "model": model,
                                     "wattage": 625, "customer": customer, "dcr": "DCR",
                                     "format_version": 2, "date_produced": "2026-10-17",
                                     "shift": 1, "sequence": 1 + _n[0], "state": "planned"})


def run(now=NOW, **q):
    p = ela.params(q)
    with store.conn() as (cx, cur):
        return ela.analyse(cx, cur, p, now, db.get_config(cur))


def drill(now=NOW, **q):
    p = ela.params(q)
    with store.conn() as (cx, cur):
        return ela.drill(cx, cur, p, now, db.get_config(cur))


def S(n):
    return "ICON625R12A17%05d" % n


def at(day, minute):
    """An ISO time on the day given, `minute` minutes after 06:00."""
    t = datetime.datetime(2026, 10, day, 6) + datetime.timedelta(minutes=minute)
    return t.isoformat()


# --------------------------------------------------------------------------
# 1. the window
# --------------------------------------------------------------------------

@test("each window is what its button says: its start, its bars, and the window before it")
def t_windows():
    w = ela.window_of("1h", NOW)
    assert (w["start"], w["end"], len(w["buckets"])) == (
        datetime.datetime(2026, 10, 17, 14, 35), NOW, 12), w
    assert w["prev_start"] == datetime.datetime(2026, 10, 17, 13, 35)
    w = ela.window_of("8h", NOW)
    assert w["start"] == datetime.datetime(2026, 10, 17, 8, 0) and len(w["buckets"]) == 8
    assert [b["label"] for b in w["buckets"]][:2] == ["08:00", "09:00"]
    w = ela.window_of("shift", NOW)
    assert w["start"] == datetime.datetime(2026, 10, 17, 14, 0) and len(w["buckets"]) == 2
    w = ela.window_of("prevshift", NOW)
    assert (w["start"], w["end"]) == (datetime.datetime(2026, 10, 17, 6), datetime.datetime(2026, 10, 17, 14))
    # 01:12 on the 18th is C shift of the 17th: "today" began 06:00 on the 17th
    night = datetime.datetime(2026, 10, 18, 1, 12)
    assert ela.window_of("today", night)["start"] == datetime.datetime(2026, 10, 17, 6)
    assert ela.window_of("shift", night)["start"] == datetime.datetime(2026, 10, 17, 22)
    w = ela.window_of("yesterday", NOW)
    assert (w["start"], w["end"], len(w["buckets"])) == (
        datetime.datetime(2026, 10, 16, 6), datetime.datetime(2026, 10, 17, 6), 24)
    w = ela.window_of("7d", NOW)
    assert w["start"] == datetime.datetime(2026, 10, 11, 6) and len(w["buckets"]) == 20
    assert [b["label"] for b in w["buckets"]][:4] == ["11-10 A", "11-10 B", "11-10 C", "12-10 A"]
    assert w["prev_end"] == NOW - datetime.timedelta(days=7)
    w = ela.window_of("30d", NOW)
    assert w["start"] == datetime.datetime(2026, 9, 18, 6) and len(w["buckets"]) == 30
    assert [len(ela.window_of("range", NOW, f, t)["buckets"]) for f, t in (
        ("2026-10-16", "2026-10-16"), ("2026-10-10", "2026-10-14"), ("2026-09-20", "2026-10-09"))] == \
        [24, 15, 20]
    for f, t, why in (("2026-10-14", "2026-10-10", "before"), ("2026-06-01", "2026-10-01", "at most"),
                      ("", "", "From date"), ("2026-10-20", "2026-10-21", "after now")):
        try:
            ela.window_of("range", NOW, f, t)
            raise AssertionError("range %s..%s accepted" % (f, t))
        except ela.Refused as e:
            assert why in str(e), e


# --------------------------------------------------------------------------
# 2. counting
# --------------------------------------------------------------------------

@test("an image counts once, by when it was taken: a continuous-mode pass apart and in no "
      "NG rate, a re-filed image as its latest filing, modules beside images")
def t_counting():
    fresh()
    cap("pre_el", S(1), "OK", "2026-10-17T09:10:00")
    cap("pre_el", S(2), "Cell Crack", "2026-10-17T09:20:00")
    cap("pre_el", S(2), "OK", "2026-10-17T10:20:00", unit=2)                 # reworked
    cap("pre_el", S(3), "OK", "2026-10-17T09:30:00", mode="continuous")      # the machine's
    cap("pre_el", S(4), "OK", "2026-10-17T09:40:00", mode="continuous")      # ... then filed NG
    cap("pre_el", S(4), "Grid Buring", "2026-10-17T09:40:00")
    cap("pre_el", S(5), "OK", "2026-10-17T09:50:00")                         # filed OK
    cap("pre_el", S(5), "Cell Crack", "2026-10-17T09:50:00",                 # re-filed NG later
        seen="2026-10-17T15:59:00")
    t = run(station="pre_el", window="today")["blocks"]["pre_el"]["tiles"]
    # decided: S1 OK, S2 NG + OK, S4 NG (its filing over the machine's pass), S5 NG
    # (re-filed: one image) - five, four of them first images; S3 the machine's
    # alone, counted apart
    assert (t["images"], t["first"], t["re"], t["modules"]) == (5, 4, 1, 4), t
    assert (t["ng_first"], t["ng_rate"], t["ng_images"], t["ng_rate_all"]) == (3, 75.0, 3, 60.0), t
    assert (t["continuous"], t["continuous_modules"]) == (1, 1), t


@test("a module's first image at a station is its first; every later one a re-image - one "
      "imaged before the window is a re-image in it; the headline NG rate is on first "
      "images and the re-images' NG is its own figure")
def t_first_images():
    fresh()
    cap("pre_el", S(1), "Cell Crack", "2026-10-16T20:00:00")                # yesterday
    cap("pre_el", S(1), "Cell Crack", "2026-10-17T07:00:00", unit=2)        # back, NG again
    cap("pre_el", S(1), "OK", "2026-10-17T08:00:00", unit=2)                # back, passed
    cap("pre_el", S(2), "OK", "2026-10-17T07:10:00")
    cap("pre_el", S(3), "String Short", "2026-10-17T07:20:00")
    cap("pre_el", S(4), "OK", "2026-10-17T07:30:00")
    t = run(station="pre_el", window="today")["blocks"]["pre_el"]["tiles"]
    assert (t["first"], t["ng_first"], t["ng_rate"]) == (3, 1, 33.3), t
    assert (t["re"], t["re_modules"], t["re_ng"], t["re_rate"]) == (2, 1, 1, 50.0), t
    assert t["ng_rate_all"] == 40.0, "every image, re-images in, beside the headline"
    d = drill(station="pre_el", window="today", kind="reimage_ng")
    assert [r[0] for r in d["rows"]] == [S(1)], d["rows"]
    d = drill(station="pre_el", window="today", kind="first_ng")
    assert [r[0] for r in d["rows"]] == [S(3)], d["rows"]
    assert drill(station="pre_el", window="today", kind="first")["total"] == 3
    assert drill(station="pre_el", window="today", kind="reimage")["total"] == 1


@test("machines are compared on first images: the machine that takes the rework is not made "
      "the worse one by the modules it re-images - those are beside, with their own rate")
def t_machines_first():
    fresh()
    rows = []
    for k in range(10):                     # each machine: 10 first images, 1 NG
        for unit in (1, 2):
            n = 100 * unit + k
            rows.append(("pre_el", S(n), "Cell Crack" if k == 0 else "OK", at(17, 10 + k), unit))
    for k in range(6):                      # machine 1 re-images six modules, four NG again
        rows.append(("pre_el", S(200 + k), "SR" if k < 4 else "OK", at(17, 120 + k), 1))
    with store.conn() as (cx, cur):
        for st, s, v, t, unit in rows:
            _row(cur, st, s, v, t, unit=unit)
    m = {x["key"]: x for x in run(station="pre_el", window="today")["blocks"]["pre_el"]["machines"]}
    one, two = m["pre_el|1|B"], m["pre_el|2|B"]
    assert (one["first"], one["rate"], two["first"], two["rate"]) == (10, 10.0, 10, 10.0), m
    assert (one["re"], one["re_ng"], one["re_rate"], two["re"]) == (6, 4, 66.7, 0), m


# --------------------------------------------------------------------------
# 3. rates, significance, coverage
# --------------------------------------------------------------------------

@test("a busier window is not a rise: the change is of the RATE on first images, with its z; "
      "only z >= 3 rose, 2-3 is one to watch")
def t_rates():
    fresh()
    rows = []
    # the window before (the 16th): 100 first images, 5 Grid Burning, 2 String Short
    for k in range(100):
        v = "Grid Burning" if k < 5 else "String Short" if k < 7 else "OK"
        rows.append(("pre_el", S(1000 + k), v, at(16, k)))
    # today: three times busier, Grid Burning at the same 5%; String Short 2% -> 8%;
    # Cell Crack from 0 to 2%
    for k in range(300):
        v = ("Grid Burning" if k < 15 else "String Short" if k < 39 else
             "Cell Crack" if k < 45 else "OK")
        rows.append(("pre_el", S(2000 + k), v, at(17, k)))
    caps(rows)
    b = run(station="pre_el", window="today")["blocks"]["pre_el"]
    r = {x["key"]: x for x in b["reasons"]}
    gb, ss, cc = r["DF-GRIDBURNING"], r["DF-STRSHORT"], r["DF-CELLCRACK"]
    assert (gb["modules"], gb["prev"], gb["rate"], gb["prev_rate"], gb["change"]) == \
        (15, 5, 5.0, 5.0, 0.0), gb
    assert gb["signal"] is None and gb["z"] == 0.0, "three times the images is not a rise"
    assert (ss["rate"], ss["prev_rate"], ss["signal"]) == (8.0, 2.0, "watch_up"), ss
    assert 2 <= ss["z"] < 3, ss
    assert cc["prev"] == 0 and cc["signal"] is None, cc
    t = b["tiles"]
    assert t["riser"] is None and t["watch"]["key"] == "DF-STRSHORT", (t["riser"], t["watch"])
    assert (t["prev"]["first"], t["prev"]["ng_rate"]) == (100, 7.0), t["prev"]
    # the same rise on more images is a signal
    fresh()
    rows = [("pre_el", S(1000 + k), "String Short" if k < 6 else "OK", at(16, k % 500))
            for k in range(300)]
    rows += [("pre_el", S(2000 + k), "String Short" if k < 24 else "OK", at(17, k % 500))
             for k in range(300)]
    caps(rows)
    t = run(station="pre_el", window="today")["blocks"]["pre_el"]["tiles"]
    assert t["riser"]["key"] == "DF-STRSHORT" and t["riser"]["z"] >= 3, t["riser"]
    assert t["prev"]["signal"] == "up", t["prev"]


@test("the comparison is made when the index covers at least 95% of the window before - "
      "a few minutes short still compares; much less says so, never a zero")
def t_coverage():
    fresh()
    cap("pre_el", S(90), "OK", "2026-10-16T06:20:00")                       # the index begins
    cap("pre_el", S(91), "Grid Burning", "2026-10-16T10:00:00")
    cap("pre_el", S(92), "Grid Burning", "2026-10-17T10:00:00")
    t = run(station="pre_el", window="today")["blocks"]["pre_el"]["tiles"]
    # the window before is 06:00-15:30 on the 16th: 9 h 10 min of 9 h 30 min indexed
    assert t["prev"]["complete"] is True and t["prev"]["covered"] == 96, t["prev"]
    fresh()
    cap("pre_el", S(90), "OK", "2026-10-16T07:00:00")
    cap("pre_el", S(92), "Grid Burning", "2026-10-17T10:00:00")
    b = run(station="pre_el", window="today")["blocks"]["pre_el"]
    assert b["tiles"]["prev"]["complete"] is False and b["tiles"]["prev"]["covered"] == 89
    assert b["tiles"]["prev"]["ng_rate"] is None and b["tiles"]["riser"] is None
    assert all(x["change"] is None and x["prev"] is None for x in b["reasons"]), b["reasons"]


# --------------------------------------------------------------------------
# 4. reasons and the Pareto
# --------------------------------------------------------------------------

@test("every spelling is one reason, read through the defect list; a folder the list does "
      "not have is its own reason, marked new - the Pareto busiest first, with share and "
      "cumulative %")
def t_reasons():
    fresh()
    for i, (v, ln) in enumerate((("Grid Buring", "B"), ("Grid burning", "A"), ("Grid Burning", "B"),
                                 (" NO IMEAGE", "B"), ("Lead OPEN/Cross", "A"), ("Corner Chipped", "B"),
                                 ("OK", "A"), ("OK", "B"))):
        cap("final_el", S(10 + i), v, "2026-10-17T11:%02d:00" % i, line=ln)
    r = run(station="final_el", window="today")["blocks"]["final_el"]["reasons"]
    got = [(x["label"], x["modules"], x["share"], x["cum"], x["new"]) for x in r]
    assert got[0] == ("Grid Burning", 3, 50.0, 50.0, False), got
    labels = {x["label"]: x for x in r}
    assert labels["No Image (Jig/Probe or JB)"]["key"] == "DF-NOIMAGE"
    assert labels["Lead Open/Cross"]["key"] == "DF-LEADOPENCROSS"
    assert labels["Corner Chipped"]["new"] is True and \
        labels["Corner Chipped"]["key"] == "new:corner chipped", labels["Corner Chipped"]
    assert r[-1]["cum"] == 100.0


@test("the Pareto ranks MODULES NG at their first image: a re-shot never counts twice, the "
      "re-images filed under a reason are beside it, the cumulative ends at 100% - and Other "
      "keeps its ranked place, at the top when it is the most")
def t_pareto_modules():
    fresh()
    caps([("pre_el", S(1), "Grid Burning", at(17, 1)),
          ("pre_el", S(1), "Grid Burning", at(17, 30)),            # a re-shot, same reason
          ("pre_el", S(1), "Grid Burning", at(17, 60)),
          ("pre_el", S(2), "other", at(17, 2)),
          ("pre_el", S(3), "Other", at(17, 3)),
          ("pre_el", S(4), "Cell Crack", at(17, 4)),
          ("pre_el", S(4), "SR", at(17, 50)),                      # back from rework, NG for SR
          ("pre_el", S(5), "OK", at(17, 5))])
    b = run(station="pre_el", window="today")["blocks"]["pre_el"]
    got = [(x["label"], x["modules"], x["re"], x["share"], x["cum"]) for x in b["reasons"]]
    # a tie in modules is broken by the re-images filed under it
    assert got == [("Other", 2, 0, 50.0, 50.0), ("Grid Burning", 1, 2, 25.0, 75.0),
                   ("Cell Crack", 1, 0, 25.0, 100.0), ("SR (Stringer Rework)", 0, 1, 0.0, None)], got
    t = b["tiles"]
    assert t["top"]["key"] == "DF-OTHER", "Other is a folder operators choose: it keeps its place"
    assert (t["other"], t["other_share"]) == (2, 50.0), t


@test("the change is against the window of the same length before it; the reason that rose "
      "beyond chance is named - and with no images that far back there is no comparison, "
      "not a zero")
def t_change():
    fresh()
    rows = [("pre_el", S(90), "OK", "2026-10-16T05:00:00")]                  # the index begins
    rows += [("pre_el", S(100 + k), "Grid Burning" if k < 2 else "OK", at(16, k)) for k in range(200)]
    rows += [("pre_el", S(400 + k), "Grid Burning" if k < 20 else "other" if k < 25 else "OK",
              at(17, k)) for k in range(200)]
    caps(rows)
    b = run(station="pre_el", window="today")["blocks"]["pre_el"]
    gb = [x for x in b["reasons"] if x["key"] == "DF-GRIDBURNING"][0]
    assert (gb["prev"], gb["rate"], gb["prev_rate"], gb["change"]) == (2, 10.0, 1.0, 9.0), gb
    assert b["tiles"]["riser"]["key"] == "DF-GRIDBURNING", b["tiles"]["riser"]
    assert b["tiles"]["prev"]["complete"] is True
    assert b["tiles"]["other_share"] == 20.0, b["tiles"]
    # yesterday's window before it (the 15th) is before the index began
    b = run(station="pre_el", window="yesterday")["blocks"]["pre_el"]
    assert b["tiles"]["prev"]["complete"] is False and b["tiles"]["riser"] is None
    assert all(x["change"] is None for x in b["reasons"]), b["reasons"]


# --------------------------------------------------------------------------
# 5. the trend and its control limits
# --------------------------------------------------------------------------

@test("the bars: each image in the bar of when it was taken, first-image NG split by reason, "
      "the NG rate on first images per bar, re-images and continuous-mode passes counted apart")
def t_trend():
    fresh()
    cap("pre_el", S(20), "Cell Crack", "2026-10-17T14:05:00")
    cap("pre_el", S(21), "OK", "2026-10-17T14:10:00")
    cap("pre_el", S(22), "OK", "2026-10-17T15:15:00", mode="continuous")
    cap("pre_el", S(23), "Cell Crack", "2026-10-17T15:20:00")
    cap("pre_el", S(20), "SR", "2026-10-17T15:25:00", unit=2)               # back, NG again
    d = run(station="pre_el", window="shift")
    tr = d["blocks"]["pre_el"]["trend"]
    assert [b["label"] for b in d["window"]["buckets"]] == ["14:00", "15:00"]
    assert [(x["first"], x["first_ng"], x["rate"], x["re"], x["re_ng"], x["cont"]) for x in tr] == \
        [(2, 1, 50.0, 0, 0, 0), (1, 1, 100.0, 1, 1, 1)], tr
    assert tr[0]["r"] == {"DF-CELLCRACK": 1} and tr[1]["r"] == {"DF-CELLCRACK": 1}, tr


def _bars(pairs):
    return [{"first": n, "first_ng": x} for x, n in pairs]


@test("control limits are Laney p': widened by the spread between bars the binomial does "
      "not explain, never narrower than plain p limits; a bar beyond 3 sigma is flagged, "
      "2-3 sigma above watched; too few bars give no limits")
def t_laney():
    import math
    # bars that vary far more than coin tosses: sigma-z well above 1, limits wider
    wild = _bars([(50, 1000), (150, 1000), (60, 1000), (140, 1000), (70, 1000), (130, 1000),
                  (80, 1000), (120, 1000), (100, 1000), (100, 1000)])
    c = ela._laney(wild)
    assert c["ok"] and c["center"] == 10.0 and c["sigma_z"] > 3, c
    plain3 = 3 * math.sqrt(0.1 * 0.9 / 1000) * 100
    assert c["bars"][0]["u3"] - 10.0 > 3 * plain3, "the limits are widened by sigma-z"
    assert c["flags"]["high3"] == 0, "a process that varies this much has no single outlier"
    # bars steadier than coin tosses: sigma-z below 1 is taken as 1 - plain p limits
    calm = _bars([(100, 1000)] * 9 + [(101, 1000)])
    c = ela._laney(calm)
    assert c["floored"] is True and abs(c["bars"][0]["u3"] - (10.0 + plain3)) < 0.06, c["bars"][0]
    # a spike in a steady run is flagged; one a little high is watched
    spike = _bars([(100, 1000), (95, 1000), (105, 1000), (100, 1000), (98, 1000), (102, 1000),
                   (100, 1000), (180, 1000), (100, 1000), (128, 1000)])
    c = ela._laney(spike)
    flags = [b["flag"] for b in c["bars"]]
    assert flags[7] == "high3", (c["sigma_z"], c["bars"][7])
    # a bar of fewer than MIN_N images is never flagged - no trend in 3 of 7
    small = spike[:]
    small[3] = {"first": 7, "first_ng": 6}
    assert ela._laney(small)["bars"][3]["flag"] is None
    # fewer than LANEY_MIN_BARS bars with images: no limits, and it says how many
    few = ela._laney(_bars([(10, 100)] * 5) + [{"first": 0, "first_ng": 0}] * 5)
    assert few["ok"] is False and few["bars_with_images"] == 5, few
    # on the answer: every station block carries its limits
    fresh()
    caps([("pre_el", S(k), "Cell Crack" if k % 9 == 0 else "OK", at(17, k)) for k in range(500)])
    ctl = run(station="pre_el", window="today")["blocks"]["pre_el"]["control"]
    assert ctl["ok"] is True and ctl["bars_with_images"] == 9, ctl


@test("reason by shift, by machine and by line, on first images: C shift after midnight is "
      "the factory day's, each machine its own column, a line with no folder said so")
def t_matrices():
    fresh()
    cap("pre_el", S(30), "SR", "2026-10-17T07:00:00", unit=1)
    cap("pre_el", S(31), "OK", "2026-10-17T15:00:00", unit=2)
    cap("pre_el", S(32), "SR", "2026-10-18T01:12:00", unit=2)             # C shift of the 17th
    cap("pre_el", S(32), "OK", "2026-10-18T01:40:00", unit=1)             # back after rework
    b = run(now=datetime.datetime(2026, 10, 18, 2, 0), station="pre_el", window="today")["blocks"]["pre_el"]
    sh = {s["key"]: (s["first"], s["first_ng"], s["re"], s["r"]) for s in b["shifts"]}
    assert sh == {"A": (1, 1, 0, {"DF-SR": 1}), "B": (1, 0, 0, {}),
                  "C": (1, 1, 1, {"DF-SR": 1})}, sh
    mc = {m["key"]: (m["first"], m["first_ng"], m["re"], m["label"]) for m in b["machines"]}
    assert mc == {"pre_el|1|B": (1, 1, 1, "Pre-EL primary · Line B"),
                  "pre_el|2|B": (2, 1, 0, "Pre-EL secondary · Line B")}, mc
    ln = {l["key"]: (l["first"], l["re"], l["configured"]) for l in b["lines"]}
    assert ln == {"A": (0, 0, False), "B": (3, 1, True)}, ln


@test("filters: line, shift and machine narrow the images; a reason narrows the NG side only "
      "(the images inspected stay the denominator); model and customer narrow to modules in the "
      "master - and each dropdown offers what the data holds under the OTHER filters")
def t_filters():
    fresh()
    master(S(40)); master(S(41), customer="SGMEDA"); master(S(42), model="ISEN600-G2X")
    cap("final_el", S(40), "DARK BUSBAR", "2026-10-17T09:00:00", line="A")
    cap("final_el", S(41), "Patches", "2026-10-17T09:10:00", line="B")
    cap("final_el", S(42), "OK", "2026-10-17T15:10:00", line="B")
    cap("final_el", S(43), "DARK BUSBAR", "2026-10-17T15:20:00", line="B")     # not in the master
    t = lambda **q: run(station="final_el", window="today", **q)["blocks"]["final_el"]["tiles"]
    assert (t(line="A")["images"], t(line="B")["images"]) == (1, 3)
    assert t(shift="B")["images"] == 2
    assert t(machine="final_el|1|B")["images"] == 3
    r = t(reason="DF-DARKBUSBAR")
    assert (r["images"], r["ng_first"], r["ng_rate"], r["other_share"]) == (4, 2, 50.0, None), r
    assert t(model="ISEN625-G12R")["images"] == 2
    assert t(customer="SGMEDA")["images"] == 1
    d = run(station="final_el", window="today", line="A")
    f = d["facets"]
    assert f["line"] == ["A", "B"], "its own pick must not hide its siblings"
    assert f["shift"] == ["A"], f["shift"]
    assert [r["key"] for r in f["reason"]] == ["DF-DARKBUSBAR"], f["reason"]
    assert f["model"] == ["ISEN625-G12R"]
    try:
        ela.params({"window": "today", "station": "final_el", "machine": "pre_el|1"})
        raise AssertionError("a malformed machine was accepted")
    except ela.Refused:
        pass


@test("first-pass yield: of the modules new to the station in the window, those that passed "
      "at their first image - one NG then passed is reworked, one still NG is said so; a module "
      "imaged before the window is not new")
def t_fpy():
    fresh()
    cap("pre_el", S(50), "OK", "2026-10-17T09:00:00")
    cap("pre_el", S(51), "Cell Crack", "2026-10-17T09:05:00")
    cap("pre_el", S(51), "OK", "2026-10-17T10:05:00", unit=2)
    cap("pre_el", S(52), "String Short", "2026-10-17T09:10:00")
    cap("pre_el", S(53), "Cell Crack", "2026-10-16T20:00:00")                # yesterday
    cap("pre_el", S(53), "OK", "2026-10-17T09:30:00")
    t = run(station="pre_el", window="today")["blocks"]["pre_el"]["tiles"]
    assert (t["first"], t["first_pass"], t["reworked"], t["still_ng"], t["fpy"]) == \
        (3, 1, 1, 1, 33.3), t
    # with B shift picked, the 09:05 NG is out of view: S(51) is not judged on its pass
    t = run(station="pre_el", window="today", shift="B")["blocks"]["pre_el"]["tiles"]
    assert t["first"] == 0, t


# --------------------------------------------------------------------------
# 6. all three side by side, and FQC
# --------------------------------------------------------------------------

@test("all three stations side by side, never added together - and the modules' journeys: NG "
      "at Pre-EL, reworked (how long), then what Final EL saw; NG at Final EL, and whether "
      "Pre-EL had seen it, passed it, or never imaged it")
def t_compare_flows():
    fresh()
    cap("pre_el", S(60), "Grid Burning", "2026-10-17T08:00:00")
    cap("pre_el", S(60), "OK", "2026-10-17T08:40:00", unit=2)               # 40 min rework
    cap("post_el", S(60), "OK", "2026-10-17T10:00:00", mode="continuous")
    cap("final_el", S(60), "Grid burning", "2026-10-17T13:00:00", line="B")   # same reason
    cap("pre_el", S(61), "OK", "2026-10-17T08:10:00")
    cap("final_el", S(61), "DARK BUSBAR", "2026-10-17T13:10:00", line="B")   # an escape
    cap("final_el", S(62), "Patches", "2026-10-17T13:20:00", line="A")       # never at Pre-EL
    cap("pre_el", S(63), "Cell Crack", "2026-10-17T08:20:00")                # still NG
    d = run(window="today")
    assert set(d["blocks"]) == {"pre_el", "post_el", "final_el"}
    assert "total" not in d["blocks"] and all("tiles" in b for b in d["blocks"].values())
    pre = {r["key"]: r for r in d["flows"]["pre_ng"]["rows"]}
    g = pre["DF-GRIDBURNING"]
    assert (g["modules"], g["reworked"], g["median_minutes"], g["at_final"], g["final_ng"],
            g["final_same"]) == (1, 1, 40, 1, 1, 1), g
    c = pre["DF-CELLCRACK"]
    assert (c["reworked"], c["no_final"], c["median_minutes"]) == (0, 1, None), c
    fin = {r["key"]: r for r in d["flows"]["final_ng"]["rows"]}
    assert fin["DF-GRIDBURNING"]["pre_same"] == 1
    assert fin["DF-DARKBUSBAR"]["pre_pass"] == 1
    assert fin["DF-PATCHES"]["no_pre"] == 1 and fin["DF-PATCHES"]["no_pre_lines"] == {"A": 1}
    assert d["flows"]["final_ng"]["total"] == {"modules": 3, "pre_same": 1, "pre_other": 0,
                                               "pre_pass": 1, "no_pre": 1}
    # the one module seen at all three counts once at EACH - no station total anywhere
    assert [d["blocks"][s]["tiles"]["modules"] for s in ("pre_el", "final_el")] == [3, 3]


def fqc(serial, outcome, el, when, quality=None, **kw):
    with store.conn() as (cx, cur):
        row = {"serial": serial, "outcome": outcome, "grade": "A" if outcome == "pass" else None,
               "mode": "confirmed", "decided_by": "fqc1", "at": when, "el_verdict": el,
               "quality_grade": quality}
        row.update(kw)
        store.insert(cur, "fqc_record", row)


@test("FQC's decision by the EL reason it read: live decisions made in the window, how many "
      "it passed and rejected, what Quality made of the rejects - superseded, cancelled and "
      "out-of-window decisions left out; each count opens its modules")
def t_fqc_table():
    fresh()
    master(S(1)); master(S(2), model="ISEN600-G2X")
    fqc(S(1), "pass", "DARK BUSBAR", "2026-10-17T09:00:00")
    fqc(S(2), "pass", "Dark Busbar", "2026-10-17T09:10:00")
    fqc(S(3), "reject", "DARK BUSBAR", "2026-10-17T09:20:00", quality="GY")
    fqc(S(4), "reject", "Patches", "2026-10-17T09:30:00")                   # Quality pending
    fqc(S(5), "pass", "OK", "2026-10-17T09:40:00")
    fqc(S(6), "pass", None, "2026-10-17T09:50:00")                          # no EL reading
    fqc(S(7), "reject", "Patches", "2026-10-17T10:00:00", superseded_by=1)  # retested
    fqc(S(8), "reject", "Patches", "2026-10-17T10:10:00", status="cancelled")
    fqc(S(9), "reject", "Patches", "2026-10-16T10:00:00")                   # yesterday
    d = run(window="today", station="final_el")
    rows = [(r["key"], r["decisions"], r["passed"], r["rejected"], r["pass_rate"], r["quality"])
            for r in d["fqc"]["rows"]]
    assert rows == [
        ("DF-DARKBUSBAR", 3, 2, 1, 66.7, {"A": 0, "GY": 1, "BGY": 0, "pending": 0}),
        ("DF-PATCHES", 1, 0, 1, 0.0, {"A": 0, "GY": 0, "BGY": 0, "pending": 1}),
        ("__clean", 1, 1, 0, 100.0, {"A": 0, "GY": 0, "BGY": 0, "pending": 0}),
        ("__none", 1, 1, 0, 100.0, {"A": 0, "GY": 0, "BGY": 0, "pending": 0})], rows
    assert d["fqc"]["decisions"] == 6
    assert "fqc" in run(window="today") and "fqc" not in run(window="today", station="pre_el")
    assert "fqc" not in run(window="today", flows="0"), "a dashboard card does not ask for it"
    m = run(window="today", station="final_el", model="ISEN600-G2X")["fqc"]["rows"]
    assert [(r["key"], r["decisions"]) for r in m] == [("DF-DARKBUSBAR", 1)], m
    dd = drill(window="today", station="final_el", kind="fqc", reason="DF-DARKBUSBAR",
               outcome="pass")
    assert sorted(r[0] for r in dd["rows"]) == [S(1), S(2)], dd["rows"]
    assert dd["columns"][2] == "FQC decision" and dd["rows"][0][2] == "Pass"
    dd = drill(window="today", station="final_el", kind="fqc", reason="DF-DARKBUSBAR",
               outcome="reject")
    assert [(r[0], r[-1]) for r in dd["rows"]] == [(S(3), "GY")], dd["rows"]
    try:
        ela.params({"window": "today", "outcome": "maybe"})
        raise AssertionError("an unknown outcome was accepted")
    except ela.Refused:
        pass


# --------------------------------------------------------------------------
# 7. every figure opens its modules
# --------------------------------------------------------------------------

@test("a dashboard card asks without the journeys across stations (flows=0) - the same "
      "figures, read quicker")
def t_no_flows():
    fresh()
    cap("pre_el", S(80), "SR", "2026-10-17T09:00:00")
    cap("final_el", S(80), "SR", "2026-10-17T13:00:00")
    with_f, without = run(window="today"), run(window="today", flows="0")
    assert "flows" in with_f and "flows" not in without
    assert with_f["blocks"] == without["blocks"]


@test("a bar, a reason, a cell or a journey lists its modules, one row a module, newest first")
def t_drill():
    fresh()
    master(S(70))
    cap("pre_el", S(70), "Cell Crack", "2026-10-17T14:05:00")
    cap("pre_el", S(71), "Cell Crack", "2026-10-17T15:05:00")
    cap("pre_el", S(72), "SR", "2026-10-17T15:10:00")
    cap("pre_el", S(73), "OK", "2026-10-17T15:15:00", mode="continuous")
    d = drill(station="pre_el", window="shift", reason="DF-CELLCRACK")
    assert d["total"] == 2 and [r[0] for r in d["rows"]] == [S(71), S(70)], d["rows"]
    assert d["rows"][1][1] == "ISEN625-G12R" and d["rows"][0][1] == "not in the master"
    assert d["columns"][:2] == ["Serial", "Model"]
    d = drill(station="pre_el", window="shift", bucket="1")
    assert sorted(r[0] for r in d["rows"]) == [S(71), S(72)], d["rows"]
    assert d["window"]["bucket"].endswith("15:00 to 16:00"), d["window"]
    d = drill(station="pre_el", window="shift", kind="first_ng", bucket="1", reason="DF-SR")
    assert [r[0] for r in d["rows"]] == [S(72)], d["rows"]
    d = drill(station="pre_el", window="shift", kind="continuous")
    assert [r[0] for r in d["rows"]] == [S(73)]
    d = drill(station="pre_el", window="shift", kind="still_ng")
    assert sorted(r[0] for r in d["rows"]) == [S(70), S(71), S(72)]
    try:
        drill(station="pre_el", window="shift", bucket="9")
        raise AssertionError("a bar past the window was accepted")
    except ela.Refused:
        pass
    cap("final_el", S(71), "Cell Crack", "2026-10-17T15:25:00")
    d = drill(window="shift", flow="final_ng", sub="pre_same")
    assert [r[0] for r in d["rows"]] == [S(71)] and d["title"].startswith("NG at Final EL"), d


# --------------------------------------------------------------------------
# 8. who sees it, and the Excel
# --------------------------------------------------------------------------

@test("EL Analysis opens for Production Incharge, Admin and Super Admin; another role is "
      "refused by the API and the screen's markup; a bad window is refused with its reason")
def t_gate():
    fresh()
    for role, want in (("Production Incharge", 200), ("Admin", 200), ("Super Admin", 200),
                       ("Quality", 403), ("FQC Operator", 403)):
        c = APP.app.test_client()
        AUTH.test_login(c, role=role)
        assert c.get("/api/el/analysis?window=today").status_code == want, role
        assert c.get("/api/el/analysis/modules?window=today").status_code == want, role
        assert c.get("/view/el-analysis").status_code == want, role
    c = APP.app.test_client()
    assert c.get("/api/el/analysis").status_code == 401
    AUTH.test_login(c)
    r = c.get("/api/el/analysis?window=fortnight")
    assert r.status_code == 400 and "No such window" in r.get_json()["why"]
    r = c.get("/api/el/analysis?window=today&station=pre_el")
    assert r.status_code == 200 and r.get_json()["ok"] and "alerts" in r.get_json()


@test("accounts made before EL Analysis existed are given it once, by role")
def t_later_screen():
    import icon_auth
    fresh()
    who = {r: AUTH.make_user(r)["login_id"] for r in ("Production Incharge", "Quality")}
    with store.conn() as (cx, cur):
        cur.execute("DELETE FROM user_screen_perm WHERE screen_id='elan'")
        got = icon_auth.grant_later_screens(cur)
        perms = {r: icon_auth.get_screen_perms(cur, l)["elan"] for r, l in who.items()}
    assert got == {who["Production Incharge"]: ["elan"]}, got
    assert perms == {"Production Incharge": {"view": True, "write": False},
                     "Quality": {"view": False, "write": False}}, perms


@test("the Excel prints the window and filters above the table: a preface, a blank row, the "
      "header, frozen and filtered from there")
def t_export_preface():
    from openpyxl import load_workbook
    fresh()
    c = APP.app.test_client()
    AUTH.test_login(c)
    r = c.post("/api/export/xlsx", json={"name": "el_analysis", "sheets": [{
        "title": "Pre-EL Pareto", "preface": ["EL Analysis · Pre-EL · Today", "Filters: none"],
        "columns": ["Reason", "NG images"], "rows": [["Grid Burning", "3"], ["=cmd", "1"]]}]})
    assert r.status_code == 200, r.get_data(as_text=True)[:200]
    ws = load_workbook(io.BytesIO(r.data))["Pre-EL Pareto"]
    vals = [[c.value for c in row] for row in ws.iter_rows(min_row=1, max_row=6)]
    assert vals[0][0] == "EL Analysis · Pre-EL · Today" and vals[1][0] == "Filters: none", vals
    assert vals[2][0] is None and vals[3][:2] == ["Reason", "NG images"], vals
    assert ws["A6"].data_type == "s", "a formula slipped through"
    assert ws.freeze_panes == "A5" and ws.auto_filter.ref == "A4:B6", (ws.freeze_panes,
                                                                        ws.auto_filter.ref)
    # no preface: exactly as before
    r = c.post("/api/export/xlsx", json={"sheets": [{"title": "T", "columns": ["a"],
                                                     "rows": [["1"]]}]})
    ws = load_workbook(io.BytesIO(r.data))["T"]
    assert ws["A1"].value == "a" and ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:A2"


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    width = max(len(n) for n, _ in _results)
    passed = failed = 0
    try:
        for name, fn in _results:
            try:
                fn()
                print("  PASS  %-*s" % (width, name))
                passed += 1
            except Exception as e:
                print("  FAIL  %-*s  %s" % (width, name, e))
                if "-v" in sys.argv:
                    traceback.print_exc()
                failed += 1
        print("\n%d passed, %d failed" % (passed, failed))
    finally:
        try:
            store.wipe()
        except Exception:
            pass
        shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(1 if failed else 0)
