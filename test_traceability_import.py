"""
ICON TRACE - the monthly traceability importer (icon_traceability_import +
the /api/prodentry/import routes).

    python test_traceability_import.py

THE RULES THIS FILE DEFENDS

  1. The workbook parses into ranges grouped by DATE then SHIFT, only the
     dates and shifts actually present - the cascade the screen offers.
  2. Every range is validated: quantity must equal the serial span, start and
     end must be one printed batch, and the row's stated wattage must equal
     the wattage in the serial (the nameplate). A row that fails is a named
     PROBLEM, never a silent import.
  3. Merged context cells (date/shift/wattage/customer) are carried down.
  4. Customers resolve through the master, '(SGS) AGRAWAL CHANNEL' included;
     an unresolved name is flagged, not blocked.
  5. CLAIM mode records a range Planning already issued, exactly as the manual
     Production Entry does, and skips a range that was never planned.
  6. BACKFILL mode creates the whole chain (indent line -> allocation -> serial
     rows, produced -> production entry) for a range this system never had,
     keeps the file's BOM as the allocation's material set, and refuses to
     duplicate serials that already exist.
  7. The real September file parses to its real shape, and its real
     data-quality faults are surfaced as problems.

Each test names the rule it defends, so a failure says which decision broke.
"""

import io, os, sys, tempfile, traceback, datetime

TMP = tempfile.mkdtemp(prefix="icontrace_trace_")
os.environ["ICON_DB_FILE"] = os.path.join(TMP, "test.db")
# The file's dates are fixed (September 2026). Claim mode keeps the 30-day
# backdate guard (backfill bypasses it, by design), which is not what this
# suite tests - without this the suite fails once September is 31 days old.
os.environ["ICON_PROD_BACKDATE_DAYS"] = "3650"

import openpyxl                                                # noqa: E402
import db                                                     # noqa: E402
import store                                                  # noqa: E402
import app as APP                                             # noqa: E402
import auth_test_helper as AUTH                               # noqa: E402
import icon_traceability_import as T                          # noqa: E402
import icon_customers as customers                            # noqa: E402

# the master's own canonical name for stock - the file calls it "SR MODULE" /
# "NORMAL" / "ICON STOCK", all of which land here
STOCK_NAME = customers.resolve("ICON Stock")["name"]

_results = []


def test(name):
    def deco(fn):
        _results.append((name, fn))
        return fn
    return deco


# --- build a workbook in memory, in the real file's shape --------------------

HEADERS = ["Date", "Shift", "Module Wattage", "Serial No.", "Ending No",
           "Quantity", "Special Customer", "Remark", "Cell Make",
           "Cell Invoice/Batch no", "Cell Type", "Bus Bar",
           "Modules Sizes H & F", "Solar Glass (Front)",
           "Solar Glass (Front) Invoice/Batch no", "Solar Glass (Back)",
           "Solar Glass (Back) Invoice/Batch no", "Cells Inter connector",
           "Ribbon Invoice/Batch no"]

BOM = ["LIONSOLAR  25.7% (210*182.2) G12R CELL", "ID20260701", "G12R", "16 BB",
       "790, 1400, 1094 MM", "KIBING (2376*1128*2 MM)", "KBM-260408",
       "KIBING (2376*1128*2 MM)", "KBM-260408", "GEBA  0.26MM", "113166"]


def wb_bytes(rows, title_block=True, start_col=1):
    """rows: list of dicts (date, shift, watt, start, end, qty, customer,
    optional bom=True). start_col mimics the real file's empty column A."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TRACEABILITY TEST"
    r = 1
    if title_block:
        ws.cell(2, 7, "ICON SOLAR-EN POWER TECHNOLOGIES PVT. LTD.  -  (TRACEABILITY)")
        ws.cell(3, 7, "Shiftwise ... for the month TEST - 2026")
        r = 12
    for c, h in enumerate(HEADERS):
        ws.cell(r, start_col + 1 + c, h)
    r += 1
    for row in rows:
        base = [row.get("date"), row.get("shift"), row.get("watt"),
                row.get("start"), row.get("end"), row.get("qty"),
                row.get("customer"), row.get("remark", True)]
        vals = base + (BOM if row.get("bom") else [])
        for c, v in enumerate(vals):
            if v is not None:
                ws.cell(r, start_col + 1 + c, v)
        r += 1
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


D1 = datetime.datetime(2026, 9, 1)
D2 = datetime.datetime(2026, 9, 2)
# format-v2 serials (10 digits after R, 4-digit running number)
def S(seq, watt=625, batch="ICON%dR129014" % 625):
    return "%s%04d" % (batch, seq)


def client():
    store.wipe()
    AUTH.ensure_auth_schema()
    c = APP.app.test_client()
    AUTH.test_login(c)
    return c


def parse_upload(c, data, name="trace.xlsx"):
    return c.post("/api/prodentry/import/parse",
                  data={"file": (io.BytesIO(data), name)},
                  content_type="multipart/form-data")


def plant_planned(serials, model="ISEN625-G12R", watt=625, customer="ICON Stock"):
    """Serial rows as Planning would have issued them (state 'planned')."""
    import icon_challan_import as CI
    with store.conn() as (cx, cur):
        iid = store.insert(cur, "indent", {"indent_no": "T/PLAN", "indent_date": "2026-09-01",
                                           "customer": customer, "created_by": "t"})
        lid = store.insert(cur, "indent_line", {
            "indent_id": iid, "line_no": 1, "model": model,
            "item_description": "SOLAR PV MODULE-%s-NDCR" % model,
            "wattage": watt, "qty": len(serials), "dcr": "NDCR"})
        aid = store.insert(cur, "allocation", {
            "indent_line_id": lid, "model": model, "wattage": watt,
            "customer": customer, "dcr": "NDCR", "date_produced": "2026-09-01",
            "shift": 1, "qty": len(serials), "seq_from": 0, "seq_to": 0,
            "created_by": "t"})
        for s in serials:
            r = CI.decompose(s)
            store.insert(cur, "serial", {
                "serial": s, "build_instance": 1, "alloc_id": aid,
                "indent_line_id": lid, "model": model, "wattage": watt,
                "customer": customer, "dcr": "NDCR",
                "format_version": r["format_version"], "date_produced": r["date_produced"],
                "shift": r["shift"], "sequence": r["sequence"], "state": "planned"})


def serial_row(s):
    with store.conn() as (cx, cur):
        r = store.one(cur, "SELECT * FROM serial WHERE serial=%s AND build_instance=1", (s,))
        return dict(r) if r else None


# --------------------------------------------------------------------------
# 1-4  parsing
# --------------------------------------------------------------------------

@test("the workbook parses into ranges grouped by date then shift - only the "
      "dates and shifts actually present")
def t_parse_cascade():
    data = wb_bytes([
        {"date": D1, "shift": "A", "watt": "625W", "start": S(1), "end": S(7), "qty": 7, "customer": "BOROSIL", "bom": True},
        {"date": D1, "shift": "A", "watt": "625W", "start": S(8), "end": S(20), "qty": 13, "customer": "BOROSIL"},
        {"date": D1, "shift": "B", "watt": "625W", "start": S(21), "end": S(30), "qty": 10, "customer": "BOROSIL"},
        {"date": D2, "shift": "A", "watt": "625W", "start": S(31), "end": S(40), "qty": 10, "customer": "BOROSIL"},
    ])
    r = T.parse(data)
    assert r["ok"], r
    assert r["dates"] == ["2026-09-01", "2026-09-02"], r["dates"]
    assert r["shifts_by_date"]["2026-09-01"] == ["A", "B"], r["shifts_by_date"]
    assert r["shifts_by_date"]["2026-09-02"] == ["A"], r["shifts_by_date"]
    assert r["summary"]["ranges"] == 4 and r["summary"]["modules"] == 40, r["summary"]
    first = r["ranges"][0]
    assert first["model"] == "ISEN625-G12R" and first["qty"] == 7, first
    assert first["customer_name"] == "Borosil Renewables Limited" and first["customer_resolved"], first


@test("a quantity that does not match the serial span is a problem, not an "
      "import")
def t_parse_qty_mismatch():
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W",
                      "start": S(1), "end": S(7), "qty": 99, "customer": "BOROSIL"}])
    r = T.parse(data)
    assert r["ranges"] == [], r["ranges"]
    assert len(r["problems"]) == 1 and "does not match the serial span" in r["problems"][0]["why"], r["problems"]


@test("a row whose stated wattage differs from the serial's is refused - the "
      "wattage in the serial is the nameplate")
def t_parse_nameplate():
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "630W",   # serial is 625
                      "start": S(1), "end": S(7), "qty": 7, "customer": "BOROSIL"}])
    r = T.parse(data)
    assert r["ranges"] == [], r["ranges"]
    assert "nameplate" in r["problems"][0]["why"], r["problems"]


@test("a malformed serial is a named problem; good rows in the same file still "
      "come through")
def t_parse_malformed():
    data = wb_bytes([
        {"date": D1, "shift": "A", "watt": "625W", "start": "ICON625R12901400050", "end": S(7), "qty": 7, "customer": "BOROSIL"},
        {"date": D1, "shift": "A", "watt": "625W", "start": S(8), "end": S(20), "qty": 13, "customer": "BOROSIL"},
    ])
    r = T.parse(data)
    assert len(r["ranges"]) == 1 and r["ranges"][0]["qty"] == 13, r["ranges"]
    assert len(r["problems"]) == 1, r["problems"]


@test("merged context cells (date, shift, customer) are carried down to the "
      "rows under them")
def t_parse_merged_ffill():
    data = wb_bytes([
        {"date": D1, "shift": "A", "watt": "625W", "start": S(1), "end": S(7), "qty": 7, "customer": "BOROSIL"},
        {"date": None, "shift": None, "watt": "625W", "start": S(8), "end": S(20), "qty": 13, "customer": None},
    ])
    r = T.parse(data)
    assert len(r["ranges"]) == 2, r["ranges"]
    assert r["ranges"][1]["date"] == "2026-09-01" and r["ranges"][1]["shift"] == "A", r["ranges"][1]
    assert r["ranges"][1]["customer_name"] == "Borosil Renewables Limited", r["ranges"][1]


@test("'(SGS) AGRAWAL CHANNEL' resolves by stripping the parenthetical tag; a "
      "genuinely unknown name is flagged")
def t_parse_customer_resolution():
    data = wb_bytes([
        {"date": D1, "shift": "A", "watt": "625W", "start": S(1), "end": S(7), "qty": 7, "customer": "(SGS) AGRAWAL CHANNEL"},
        {"date": D1, "shift": "A", "watt": "625W", "start": S(8), "end": S(20), "qty": 13, "customer": "ZZZ TRANSPORT CO"},
    ])
    r = T.parse(data)
    assert r["ranges"][0]["customer_resolved"] and "Agrawal" in r["ranges"][0]["customer_name"], r["ranges"][0]
    assert not r["ranges"][1]["customer_resolved"], r["ranges"][1]
    assert "ZZZ TRANSPORT CO" in r["unresolved_customers"], r["unresolved_customers"]


@test("'SR MODULE' is not a customer - it is a String Rework module: it resolves "
      "to ICON Stock, is marked rework, and is NOT flagged as unknown; 'NORMAL' "
      "is ordinary ICON Stock (not rework)")
def t_parse_rework_category():
    data = wb_bytes([
        {"date": D1, "shift": "A", "watt": "625W", "start": S(1), "end": S(7), "qty": 7, "customer": "SR MODULE"},
        {"date": D1, "shift": "A", "watt": "625W", "start": S(8), "end": S(20), "qty": 13, "customer": "NORMAL"},
    ])
    r = T.parse(data)
    sr, normal = r["ranges"][0], r["ranges"][1]
    assert sr["customer_resolved"] and sr["customer_name"] == STOCK_NAME and sr["rework"] is True, sr
    assert normal["customer_resolved"] and normal["customer_name"] == STOCK_NAME and normal["rework"] is False, normal
    assert r["unresolved_customers"] == [], r["unresolved_customers"]


@test("the BOM parses into allocation-material rows, with cell efficiency "
      "pulled out of the cell-make text")
def t_parse_bom():
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "BOROSIL", "bom": True}])
    r = T.parse(data)
    mats = T.bom_materials(r["ranges"][0]["bom"])
    by_no = {m["material_no"]: m for m in mats}
    assert 5 in by_no and by_no[5]["efficiency"] == "25.7%", by_no          # G12R cell
    assert by_no[5]["batch"] == "ID20260701", by_no
    assert 6 in by_no and by_no[6]["vendor"].startswith("KIBING"), by_no    # glass front
    assert 12 in by_no and by_no[12]["batch"] == "113166", by_no            # ribbon


# --------------------------------------------------------------------------
# 5  claim mode
# --------------------------------------------------------------------------

@test("CLAIM mode records a range Planning already issued: serials move to "
      "produced and a production entry is written")
def t_apply_claim():
    c = client()
    serials = [S(i) for i in range(1, 8)]
    plant_planned(serials)
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "ICON STOCK"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    assert rng["all_present"], rng
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "TEST INCHARGE"}).get_json()
    assert r["recorded"] == 1 and r["results"][0]["action"] == "claimed", r
    assert serial_row(S(1))["state"] == "produced" and serial_row(S(1))["prod_entry_id"], serial_row(S(1))
    with store.conn() as (cx, cur):
        assert store.one(cur, "SELECT COUNT(*) AS n FROM production_entry")["n"] == 1


@test("CLAIM mode skips a range this system never planned, and says to use "
      "backfill")
def t_apply_claim_unplanned():
    c = client()
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "BOROSIL"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    assert not rng["all_present"] and rng["in_system"] == 0, rng
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "X"}).get_json()
    assert r["recorded"] == 0 and r["results"][0]["action"] == "skipped", r
    assert "backfill" in r["results"][0]["why"].lower(), r


# --------------------------------------------------------------------------
# 6-8  backfill mode
# --------------------------------------------------------------------------

@test("BACKFILL mode builds the whole chain for an unplanned range: indent "
      "line, allocation, serials (produced), production entry - and keeps the "
      "BOM as the allocation's material set")
def t_apply_backfill():
    c = client()
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "BOROSIL", "bom": True}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "NIGHT INCHARGE",
                     "backfill": True, "dcr": "NDCR"}).get_json()
    res = r["results"][0]
    assert r["recorded"] == 1 and res["action"] == "backfilled", r
    # 7 serials, produced, nameplate matches the allocation
    s1 = serial_row(S(1))
    assert s1["state"] == "produced" and s1["wattage"] == 625, s1
    assert s1["customer"] == "Borosil Renewables Limited" and s1["dcr"] == "NDCR", s1
    with store.conn() as (cx, cur):
        assert store.one(cur, "SELECT COUNT(*) AS n FROM serial")["n"] == 7
        assert store.one(cur, "SELECT COUNT(*) AS n FROM production_entry")["n"] == 1
        aid = s1["alloc_id"]
        assert aid == res["alloc_id"]
        mats = store.rows(cur, "SELECT material_no, vendor, efficiency, batch "
                               "FROM allocation_material WHERE alloc_id=%s", (aid,))
        by_no = {m["material_no"]: m for m in mats}
        assert 5 in by_no and by_no[5]["efficiency"] == "25.7%", by_no
        assert 12 in by_no and by_no[12]["batch"] == "113166", by_no
        # the backfill indent chain exists and is reusable
        ind = store.one(cur, "SELECT indent_no, customer FROM indent WHERE indent_no=%s",
                        ("BACKFILL/C0002",))
        assert ind and ind["customer"] == "Borosil Renewables Limited", ind


@test("BACKFILL against an unresolved customer imports under ICON Stock and "
      "flags it")
def t_apply_backfill_unresolved_customer():
    c = client()
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "ZZZ TRANSPORT CO"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "X", "backfill": True}).get_json()
    res = r["results"][0]
    assert res["action"] == "backfilled" and res["customer_flagged"], res
    assert serial_row(S(1))["customer"] == STOCK_NAME, serial_row(S(1))


@test("BACKFILL of a String Rework (SR MODULE) range: ICON Stock, and every "
      "serial carries the rework flag")
def t_apply_backfill_rework():
    c = client()
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "SR MODULE"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    assert rng["rework"] is True and not rng["customer_resolved"] is False, rng
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "X", "backfill": True}).get_json()
    assert r["results"][0]["action"] == "backfilled" and r["results"][0]["rework"], r
    s1 = serial_row(S(1))
    assert s1["customer"] == STOCK_NAME and s1["rework"] == 1, s1


@test("a module FQC judged BEFORE the backfill created it carries on from that "
      "decision - a pass is graded A and packs, a reject goes to Quality - not "
      "left 'produced' (637 modules on 01-10-2026)")
def t_backfill_carries_standing_fqc():
    c = client()
    with store.conn() as (cx, cur):
        store.insert(cur, "fqc_record", {"serial": S(1), "outcome": "pass",
            "grade": "A", "mode": "confirmed", "decided_by": "t",
            "at": "2026-09-01T07:00:00"})
        store.insert(cur, "fqc_record", {"serial": S(2), "outcome": "reject",
            "mode": "confirmed", "decided_by": "t", "defect": "Cell Crack",
            "at": "2026-09-01T07:01:00"})
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "BOROSIL"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "X", "backfill": True}).get_json()
    assert r["results"][0]["action"] == "backfilled", r
    assert r["results"][0]["fqc_applied"] == 2, r
    assert (serial_row(S(1))["state"], serial_row(S(1))["grade"]) == ("graded", "A")
    assert serial_row(S(2))["state"] == "rejected", serial_row(S(2))
    assert serial_row(S(3))["state"] == "produced", "an uninspected module stays produced"
    # and Packing takes the passed one
    box = c.post("/api/box/open", json={"grade": "A", "model": "ISEN625-G12R",
                                        "capacity": 36}).get_json()
    r = c.post("/api/box/%d/scan" % box["box_id"], json={"serial": S(1)}).get_json()
    assert r["ok"], r


@test("a String Rework module warns at packing (soft confirm) and packs once "
      "confirmed - it is ICON Stock, mixable with regular modules")
def t_pack_rework_warns():
    c = client()
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "SR MODULE"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    c.post("/api/prodentry/import/apply",
           json={"ranges": [rng], "incharge": "X", "backfill": True})
    # make one pass FQC so it is packable (backfill leaves it 'produced') -
    # with the FQC record a pass leaves, which Packing reads
    with store.conn() as (cx, cur):
        store.insert(cur, "fqc_record", {"serial": S(1), "outcome": "pass",
            "grade": "A", "mode": "confirmed", "decided_by": "t",
            "at": "2026-09-01T09:00:00"})
        cur.execute("UPDATE serial SET state='graded', grade='A' WHERE serial=%s", (S(1),))
    box = c.post("/api/box/open", json={"grade": "A", "model": "ISEN625-G12R",
                                        "capacity": 36}).get_json()
    # first scan: a soft confirm, not a refusal, and nothing packed yet
    r = c.post("/api/box/%d/scan" % box["box_id"], json={"serial": S(1)}).get_json()
    assert r["ok"] is False and r.get("confirm") == "rework", r
    assert "String Rework" in r["why"], r
    assert serial_row(S(1))["state"] == "graded", "not packed before confirming"
    # confirmed: it packs
    r = c.post("/api/box/%d/scan" % box["box_id"],
               json={"serial": S(1), "confirm_rework": True}).get_json()
    assert r["ok"] and serial_row(S(1))["state"] == "packed", r
    # box/check surfaces the rework flag for the preview
    chk = c.get("/api/box/check?serial=%s&grade=A" % S(2)).get_json()
    assert chk["rework"] is True, chk


@test("BACKFILL of a range already fully in the system creates nothing - it "
      "never duplicates what is there")
def t_apply_backfill_no_duplicate():
    c = client()
    plant_planned([S(i) for i in range(1, 8)])
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "BOROSIL"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "X", "backfill": True}).get_json()
    assert r["recorded"] == 0 and r["results"][0]["action"] == "skipped", r
    with store.conn() as (cx, cur):
        assert store.one(cur, "SELECT COUNT(*) AS n FROM serial")["n"] == 7   # unchanged


@test("BACKFILL of a PARTIAL range imports only the remaining serials, leaving "
      "the ones already present untouched - '276/377' fills the other 101")
def t_apply_backfill_partial_remaining():
    c = client()
    # 3 of the 7 already planned (a gap: 1,2,3 present; 4-7 missing)
    plant_planned([S(1), S(2), S(3)])
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "BOROSIL"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    assert rng["in_system"] == 3 and not rng["all_present"], rng
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "X", "backfill": True}).get_json()
    res = r["results"][0]
    assert res["action"] == "backfilled" and res["qty"] == 4, res   # only the 4 missing
    assert res["skipped_present"] == 3, res
    # all 7 now exist; the first 3 are still 'planned' (untouched), 4-7 produced
    with store.conn() as (cx, cur):
        assert store.one(cur, "SELECT COUNT(*) AS n FROM serial")["n"] == 7
    assert serial_row(S(1))["state"] == "planned", "the present ones were not touched"
    assert serial_row(S(4))["state"] == "produced" and serial_row(S(7))["state"] == "produced"


@test("import remaining handles a GAP in the middle - two runs become two "
      "production entries, each one contiguous")
def t_apply_backfill_gap_two_runs():
    c = client()
    plant_planned([S(4), S(5)])          # middle present; 1-3 and 6-7 missing
    data = wb_bytes([{"date": D1, "shift": "A", "watt": "625W", "start": S(1),
                      "end": S(7), "qty": 7, "customer": "BOROSIL"}])
    rng = parse_upload(c, data).get_json()["ranges"][0]
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": [rng], "incharge": "X", "backfill": True}).get_json()
    res = r["results"][0]
    assert res["action"] == "backfilled" and res["qty"] == 5 and res["entries"] == 2, res
    with store.conn() as (cx, cur):
        pe = store.rows(cur, "SELECT start_serial, end_serial, qty FROM production_entry "
                             "ORDER BY entry_id")
    spans = sorted((p["start_serial"], p["end_serial"], p["qty"]) for p in pe)
    assert spans == [(S(1), S(3), 3), (S(6), S(7), 2)], spans


@test("one failing range does not undo the others in the same import - each "
      "is applied under its own savepoint")
def t_apply_isolates_failures():
    c = client()
    data = wb_bytes([
        {"date": D1, "shift": "A", "watt": "625W", "start": S(1), "end": S(7), "qty": 7, "customer": "BOROSIL"},
        {"date": D1, "shift": "A", "watt": "625W", "start": S(8), "end": S(20), "qty": 13, "customer": "BOROSIL"},
    ])
    ranges = parse_upload(c, data).get_json()["ranges"]
    plant_planned([S(i) for i in range(1, 8)])     # only the FIRST range exists -> its backfill skips
    r = c.post("/api/prodentry/import/apply",
               json={"ranges": ranges, "incharge": "X", "backfill": True}).get_json()
    actions = sorted(x["action"] for x in r["results"])
    assert actions == ["backfilled", "skipped"], actions
    assert serial_row(S(8)) and serial_row(S(8))["state"] == "produced", "the good range still landed"


# --------------------------------------------------------------------------
# 7  the real file
# --------------------------------------------------------------------------

REAL = ("C:/Users/X/.claude/uploads/36f75a19-12fd-48d0-a090-aadf49a67a63/"
        "775c1e29-09___UNIT-2_TRACEABILITY_REPORT__SEPTEMBER_-2026-1.xlsx")


@test("the real September file parses to its real shape, and its real "
      "data-quality faults surface as problems (not silent imports)")
def t_real_file():
    if not os.path.exists(REAL):
        print("      (real file not present - skipped)")
        return
    r = T.parse(REAL)
    assert r["ok"] and r["month_label"] == "SEPTEMBER - 2026", r.get("month_label")
    assert r["summary"]["ranges"] == 260, r["summary"]
    assert len(r["dates"]) == 29 and r["dates"][0] == "2026-09-01", r["dates"][:2]
    assert r["shifts_by_date"]["2026-09-01"] == ["A", "B", "C"], r["shifts_by_date"]["2026-09-01"]
    # the four malformed serials I found by hand are all flagged
    assert r["summary"]["problems"] == 4, r["problems"]
    bad = " ".join(p["why"] for p in r["problems"])
    assert "ICON625R12912400050" in bad, bad
    # SR MODULE (rework) and NORMAL (stock) are categories, not unknown
    # customers - only a real near-miss alias is left flagged
    assert "SR MODULE" not in r["unresolved_customers"], r["unresolved_customers"]
    assert "AGNI GREEN" in r["unresolved_customers"], r["unresolved_customers"]
    assert any(x["rework"] for x in r["ranges"]), "the file's SR MODULE ranges are marked rework"


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    width = max(len(n) for n, _ in _results)
    passed = failed = 0
    try:
        for name, fn in _results:
            try:
                fn()
                print("  PASS  %-*s" % (width, name))
                passed += 1
            except Exception as e:
                print("  FAIL  %-*s  %s" % (width, name, e))
                if "-v" in sys.argv:
                    traceback.print_exc()
                failed += 1
        print("\n%d passed, %d failed" % (passed, failed))
    finally:
        import shutil
        shutil.rmtree(TMP, ignore_errors=True)
    sys.exit(1 if failed else 0)
