"""
ICON TRACE - the traceability report as Excel.

    python test_trace_export.py          (the last test needs Playwright)

Mukesh: "Production export in this format, month wise data, range can be selected
from date to date and only single date is selected for 1 day data" and "we don't
merge 2 materials in a single column like both potting materials, etc - single
material each column."

THE RULES THIS FILE DEFENDS

  1. THE PERIOD: a month, a range of dates, or - FROM alone - one day. Counted by
     the shift the production RAN in (prod_date), a cancelled entry never
     appears, an end before its start / no dates / over a year is refused.
  2. ONE MATERIAL, ONE COLUMN: both potting parts, the centre and the edge ribbon
     and the cell's make, efficiency and batch each have their own column; a group
     of alternatives (the 15 / 20 mm tape, the ribbon thicknesses, the frame's
     mounting holes) shares one column and says which variant was used.
  3. The report's own words: SR MODULE for rework, NORMAL for unallocated stock,
     the master's name for a customer; the wattage summary is a formula.
  4. Gated on the Production Entry screen's view flag.
  5. THE SCREEN: a month fills FROM / TO; the button downloads the file; an empty
     request is refused on screen.

ui_harness is imported FIRST: it fixes the throwaway database path.
"""

import datetime, io, os, sys, traceback

import ui_harness as H                                       # noqa: E402  (first: sets the DB path)
import db                                                    # noqa: E402
import store                                                 # noqa: E402
import auth_test_helper as AUTH                              # noqa: E402
import test_bom_match as BM                                  # noqa: E402  (its full-shape workbook)
import icon_trace_export as TX                               # noqa: E402
import openpyxl                                              # noqa: E402

_results = []


def test(name):
    def deco(fn):
        _results.append((name, fn))
        return fn
    return deco


def D(day):
    return datetime.datetime(2026, 9, day)


SHIFT_OF_DAY = {1: 1, 2: 2, 3: 3, 4: 1}          # A, B, C, A


def S(day, i):
    """ICON625R + 129 + day + shift number + 4-digit sequence (ICON625R1293030001 = 30 Sep, C, 1)."""
    return "ICON625R129%02d%d%04d" % (day, SHIFT_OF_DAY[day], i)


def world():
    """Production in the system: the 1st (A, 5 modules), the 2nd (B, 6 for Borosil
    and 3 string-rework), the 3rd (C, 4 of unallocated stock), and an entry on the
    4th that is then cancelled."""
    c = BM.client()
    rows = [dict(date=D(1), shift="A", start=S(1, 1), end=S(1, 5), qty=5),
            dict(date=D(2), shift="B", start=S(2, 1), end=S(2, 6), qty=6),
            dict(date=D(2), shift="B", start=S(2, 11), end=S(2, 13), qty=3, cust="SR MODULE"),
            dict(date=D(3), shift="C", start=S(3, 1), end=S(3, 4), qty=4, cust="NORMAL"),
            dict(date=D(4), shift="A", start=S(4, 1), end=S(4, 2), qty=2)]
    p = BM.parse_up(c, BM.workbook(rows=rows)).get_json()
    assert p["ok"] and len(p["ranges"]) == 5, p
    r = c.post("/api/prodentry/import/apply", json={"ranges": p["ranges"], "incharge": "NIGHT INCHARGE",
                                                    "backfill": True, "dcr": "NDCR"}).get_json()
    assert r["recorded"] == 5, r
    with store.conn() as (cx, cur):
        cur.execute("UPDATE production_entry SET status='cancelled' WHERE prod_date='2026-09-04'")
    return c


def sheet(c, qs):
    r = c.get("/export/traceability.xlsx" + qs)
    assert r.status_code == 200, (r.status_code, r.get_data(as_text=True)[:300])
    return openpyxl.load_workbook(io.BytesIO(r.data)).active, r


def table(ws):
    heads = [ws.cell(12, k).value for k in range(2, ws.max_column + 1)]
    out = []
    for r in range(13, ws.max_row + 1):
        if ws.cell(r, 2).value is None:
            continue
        row = {}
        for i, h in enumerate(heads):
            row.setdefault(h, ws.cell(r, 2 + i).value)      # the two "Remark" columns: the first wins
        out.append(row)
    return heads, out


@test("a MONTH, a range, and a single day: the same rows the system holds, by the shift "
      "they ran in - a cancelled entry, and days outside the period, never appear")
def t_periods():
    c = world()
    ws, r = sheet(c, "?month=2026-09")
    assert [x["Serial No."] for x in table(ws)[1]] == [S(1, 1), S(2, 1), S(2, 11), S(3, 1)], table(ws)[1]
    assert "SEPTEMBER-2026" in r.headers["Content-Disposition"]
    assert "for the month SEPTEMBER - 2026" in ws["G3"].value and ws.title == "TRACEABILITY SEP-2026"
    ws, r = sheet(c, "?from=2026-09-02&to=2026-09-03")
    assert [x["Serial No."] for x in table(ws)[1]] == [S(2, 1), S(2, 11), S(3, 1)]
    assert "02-09-2026 to 03-09-2026" in r.headers["Content-Disposition"]
    ws, r = sheet(c, "?from=2026-09-02")                     # FROM alone: ONE day
    assert [x["Serial No."] for x in table(ws)[1]] == [S(2, 1), S(2, 11)], table(ws)[1]
    assert "REPORT 02-09-2026.xlsx" in r.headers["Content-Disposition"] and ws.title == "TRACEABILITY 02.09.2026"
    assert "for 02-09-2026" in ws["G3"].value
    ws, _ = sheet(c, "?to=2026-09-03")                       # TO alone is one day too
    assert [x["Serial No."] for x in table(ws)[1]] == [S(3, 1)]
    ws, _ = sheet(c, "?from=2026-09-04")                     # only the cancelled entry
    assert table(ws)[1] == [], "a cancelled entry was exported"


@test("the period is refused, with a reason, when it cannot be meant: nothing picked, an "
      "end before its start, not a date, more than a year")
def t_refusals():
    c = world()
    for qs, word in (("", "Pick a month"), ("?from=2026-09-05&to=2026-09-01", "before"),
                     ("?from=banana", "not a date"), ("?month=2026-13", "not a date"),
                     ("?from=2025-01-01&to=2026-09-01", "at most")):
        r = c.get("/export/traceability.xlsx" + qs)
        assert r.status_code == 400 and word in r.get_json()["why"], (qs, r.status_code, r.get_data(as_text=True)[:200])


@test("ONE MATERIAL, ONE COLUMN: potting A and B, the centre and the edge ribbon, the cell's "
      "make / efficiency / batch each have their own - and a variant is named beside its make")
def t_one_column_each():
    ws, _ = sheet(world(), "?from=2026-09-01")
    heads, rows = table(ws)
    row = rows[0]
    for h in ("Potting Material — Part A", "Potting Material — Part B",
              "String Inter Connector — Centre", "String Inter Connector — Edge",
              "Cell Make", "Cell Efficiency", "Cell Invoice/Batch no",
              "Cell Alignment Tape", "Lead Bending Tape", "Edge Sealing Tape", "Junction Box 30 A",
              "EPE Strip (Output Patti)", "Back Label", "Barcode Label", "Pallet Packing", "RFID Sticker",
              "Pallet Packing Invoice/Batch no", "Potting Material — Part B Invoice/Batch no"):
        assert h in heads, h
    assert sorted(h for h in heads if heads.count(h) > 1) == ["Remark", "Remark"], "a column heading repeats"
    assert "Potting Material" not in heads and "String Inter Connector" not in heads
    assert row["Potting Material — Part A"] == "Fasto" and row["Potting Material — Part B"] == "Fasto"
    assert row["Potting Material — Part A Invoice/Batch no"] == "TIFAM2627/785"
    assert row["Cell Make"] == "Lion Solar" and row["Cell Efficiency"] == "25.7%"
    assert row["Cell Invoice/Batch no"] == "ID20260701"
    assert row["String Inter Connector — Centre"].startswith("Geba Copper"), row["String Inter Connector — Centre"]
    assert row["String Inter Connector — Edge"].endswith("4.0 x 0.42 mm"), row["String Inter Connector — Edge"]
    assert row["Lead Bending Tape"] == "H.B. Fuller · 15 mm width", row["Lead Bending Tape"]
    assert row["Back Label"] == "Kvell" and row["Barcode Label"] == "Kvell" and row["Pallet Packing"] == "Manmohan"
    assert row["Cell Type"] == "G12R" and row["Bus Bar"] == "16 BB"
    assert row["Modules Sizes H & F"] == "790, 1400, 1094 MM", row["Modules Sizes H & F"]
    assert row["Aluminium Frame"].startswith("Jiangyin Yuanshuo (YS) · 2382 x 1134 x 30 mm"), row["Aluminium Frame"]
    assert row["Shift Incharge"] == "Night Incharge" and row["Module Wattage"] == "625W"
    assert row["Date"].date() == datetime.date(2026, 9, 1) and row["Shift"] == "A" and row["Quantity"] == 5


@test("the report's own words: Borosil by its master name, SR MODULE for rework, NORMAL for "
      "unallocated stock; the check column and the wattage summary are formulas")
def t_words_and_formulas():
    ws, _ = sheet(world(), "?month=2026-09")
    _h, rows = table(ws)
    assert [x["Special Customer"] for x in rows] == ["Borosil Renewables Limited", "Borosil Renewables Limited",
                                                      "SR MODULE", "NORMAL"], [x["Special Customer"] for x in rows]
    assert ws["E5"].value == "=SUM(G13:G16)" and ws["E6"].value == '=SUMIF(D13:D16,"625W",G13:G16)'
    assert [ws["B%d" % k].value for k in range(6, 10)] == [
        "Total 625W Production", "Total 620W Production", "Total 630W Production", "Total 590W Production"]
    assert ws["I13"].value.startswith('=IF($G13=""')            # the report's quantity-vs-serials check
    assert any(str(m).startswith("G2:") for m in ws.merged_cells.ranges), "the title is not merged"
    assert ws.freeze_panes == "C13"


@test("AGNI GREEN is an alias of Agni Green Power Limited (Mz): the import files the run against "
      "that customer - not Icon Stock - and the report prints its master name")
def t_agni_green():
    c = BM.client()
    p = BM.parse_up(c, BM.workbook(rows=[dict(date=D(1), shift="A", start=S(1, 1), end=S(1, 5), qty=5,
                                              cust="AGNI GREEN")])).get_json()
    assert p["ranges"][0]["customer_resolved"] and p["unresolved_customers"] == [], p["unresolved_customers"]
    r = c.post("/api/prodentry/import/apply", json={"ranges": p["ranges"], "incharge": "NIGHT INCHARGE",
                                                    "backfill": True, "dcr": "NDCR"}).get_json()
    assert r["recorded"] == 1, r
    ws, _ = sheet(c, "?from=2026-09-01")
    assert table(ws)[1][0]["Special Customer"] == "Agni Green Power Limited (Mz)", table(ws)[1][0]
    with store.conn() as (cx, cur):
        assert store.one(cur, "SELECT customer FROM serial WHERE serial=%s", (S(1, 1),))["customer"]             .upper().startswith("AGNI GREEN POWER"), "filed against Icon Stock"


@test("the summary lists the wattages built, most first, and folds the rest into 'other'; "
      "a legacy-only material has a column only when something was recorded against it")
def t_summary_and_legacy():
    rows = [{"wattage": w, "qty": q} for w, q in ((625, 100), (600, 50), (590, 20), (620, 10), (630, 5), (605, 1))]
    shown, other = TX.summary_watts(rows)
    assert shown == [625, 600, 590] and other is True, (shown, other)
    shown, other = TX.summary_watts([{"wattage": 625, "qty": 9}])
    assert shown == [625, 620, 630, 590] and other is False, shown
    cat = [{"n": 1, "name": "A", "cat": "Cell", "series": ""},
           {"n": 2, "name": "Old", "cat": "Cell", "series": "", "legacy": True}]
    assert [g["name"] for g in TX.material_groups(cat, ["Cell"])] == ["A"]
    assert [g["name"] for g in TX.material_groups(cat, ["Cell"], recorded={2})] == ["A", "Old"]
    lab = [{"n": 24, "name": "Back Label 590WP", "cat": "L", "series": "LABEL"},
           {"n": 28, "name": "Back Label 625WP", "cat": "L", "series": "LABEL"}]
    assert [g["name"] for g in TX.material_groups(lab, ["L"])] == ["Back Label"]


@test("gated on the Production Entry screen: no session 401; a signed-in account that may view it 200")
def t_gate():
    world()
    anon = H.APP.app.test_client()
    assert anon.get("/export/traceability.xlsx?month=2026-09").status_code == 401
    c = H.APP.app.test_client()
    AUTH.test_login(c)
    assert c.get("/export/traceability.xlsx?month=2026-09").status_code == 200


@test("on the Production Entry screen: a month fills FROM and TO; the button downloads "
      "the file; FROM alone is one day; nothing picked is refused on screen")
def t_screen():
    world()
    with H.browser() as b:
        pg = H.open_page(b, "prodentry")
        pg.wait_for_selector("#peExportXlsx", timeout=8000)
        pg.fill("#peExportMonth", "2026-09")
        pg.wait_for_timeout(500)
        assert pg.input_value("#peFilterFrom") == "2026-09-01" and pg.input_value("#peFilterTo") == "2026-09-30"
        with pg.expect_download(timeout=10000) as dl:
            pg.click("#peExportXlsx")
        assert dl.value.suggested_filename == "UNIT-2 TRACEABILITY REPORT SEPTEMBER-2026.xlsx", dl.value.suggested_filename
        path = os.path.join(os.environ.get("TEMP", "."), "trace_dl.xlsx")
        dl.value.save_as(path)
        assert len(table(openpyxl.load_workbook(path).active)[1]) == 4
        pg.fill("#peExportMonth", "")
        pg.fill("#peFilterTo", "")
        pg.fill("#peFilterFrom", "2026-09-02")
        with pg.expect_download(timeout=10000) as dl:
            pg.click("#peExportXlsx")
        assert dl.value.suggested_filename == "UNIT-2 TRACEABILITY REPORT 02-09-2026.xlsx"
        pg.fill("#peFilterFrom", "")
        pg.click("#peExportXlsx")
        pg.wait_for_timeout(500)
        assert "Pick a month" in pg.inner_text("body"), "an empty request was not refused on screen"
        assert not pg.errors, pg.errors


if __name__ == "__main__":
    sys.stdout.reconfigure(errors="replace")
    only = [a for a in sys.argv[1:] if not a.startswith("-")]
    width = max(len(n) for n, _ in _results)
    passed = failed = 0
    try:
        for name, fn in _results:
            if only and not any(o in fn.__name__ for o in only):
                continue
            try:
                fn()
                print("  PASS  %-*s" % (width, name))
                passed += 1
            except Exception as e:
                print("  FAIL  %-*s  %s" % (width, name, str(e)[:500]))
                if "-v" in sys.argv:
                    traceback.print_exc()
                failed += 1
        print("\n%d passed, %d failed" % (passed, failed))
    finally:
        H.cleanup()
    sys.exit(1 if failed else 0)
